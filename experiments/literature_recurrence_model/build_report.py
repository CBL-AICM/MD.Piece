"""產生科展說明書（中文、英文各一份，.docx）。所有數字由 outputs/results.json 與 plot_data 讀入，不手抄。

格式：封面含作品名稱；內文中文新細明體、英數 Times New Roman，皆 12 pt；頁尾置中頁碼（封面不編）；
圖表依 APA 第 7 版（圖表編號粗體、標題另起一行、表格三線式、註解在下方）；參考文獻懸掛縮排。

跑法： python experiments/literature_recurrence_model/build_report.py
"""
from __future__ import annotations

import json
import os
import re
import sys

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
OUT = os.path.join(HERE, "outputs")
PD = os.path.join(OUT, "plot_data")
FIG = os.path.join(OUT, "figures")
DOCS = os.path.join(os.path.dirname(os.path.dirname(HERE)), "docs", "research", "literature_recurrence_model")
os.makedirs(DOCS, exist_ok=True)
ZH_FONT, EN_FONT, SIZE = "新細明體", "Times New Roman", Pt(12)


# ───────────── 數字 ─────────────
def numbers():
    r = json.load(open(os.path.join(OUT, "results.json"), encoding="utf-8"))
    cross = pd.read_csv(os.path.join(PD, "交叉對比_逐病一年復發率.csv"))
    tl = pd.read_csv(os.path.join(PD, "時間線_風險三分組.csv"))
    refs = pd.read_csv(os.path.join(OUT, "references_used.csv"), dtype=str).fillna("")
    lit = r["literature"]
    mc, mo = r["models"]["calibrated"], r["models"]["original"]
    me = r["mechanisms"]
    per = pd.DataFrame(mc["per_disease"])
    w = per.events
    k_lit = [c["k"] for c in lit["curves"]]
    dr_new = np.array(me["calibrated"]["dose_response_rate_y2"])
    dr_old = np.array(me["original"]["dose_response_rate_y2"])
    an = cross.dropna(subset=["lit_target"])
    tl360 = tl[tl.day == 360].set_index("組別")
    N = {
        "r": r, "cross": cross, "refs": refs, "per": per,
        "n_pubmed": int((refs.pmid != "").sum()), "n_methods": int(((refs.category == "F") & (refs.pmid != "")).sum()),
        "n_rate_est": sum(v["k"] for v in lit["targets"].values()),
        "n_rate_studies": len({p for _, p, *_ in __import__("literature").RATES}),
        "H1": lit["pooled"]["H1"], "H2": lit["pooled"]["H2"], "curves": lit["curves"],
        "k_lit_med": float(np.median(k_lit)), "k_lit_below1": int(sum(k < 1 for k in k_lit)),
        "k_orig_med": float(cross.weibull_k_orig.median()), "k_new_med": float(cross.weibull_k_calib.median()),
        "mech": me, "ms": lit["ms_prior_relapse_rr"], "dr_new": dr_new / dr_new[0], "dr_old": dr_old / dr_old[0],
        "ag": r["agreement"], "orig_min": float(an.orig_p_y2.min()), "orig_max": float(an.orig_p_y2.max()),
        "mc": mc, "mo": mo, "sens": r["sensitivity"], "prev": r["previous_study"],
        "within_rec": float(np.average(per.auc_rec730, weights=w)), "within_static": float(np.average(per.auc_static, weights=w)),
        "within_lit": float(np.average(per.auc_lit, weights=w)), "perm": mc["permutation_auc"],
        "tl": tl360, "ep_new": cross.set_index("en").calib_episode_days, "ep_old": cross.set_index("en").orig_episode_days,
        "settings": r["settings"],
    }
    return N


def f2(x):
    return f"{x:.2f}"


def f3(x):
    return f"{x:.3f}"


def ci(v, d=3):
    return f"{v[0]:.{d}f}（{v[1]:.{d}f}–{v[2]:.{d}f}）"


def ci_en(v, d=3):
    return f"{v[0]:.{d}f} ({v[1]:.{d}f}–{v[2]:.{d}f})"


def pct(x, d=1):
    return f"{100 * x:.{d}f}%"


# ───────────── docx 工具 ─────────────
def set_run_font(run, size=SIZE, bold=None, italic=None):
    run.font.name = EN_FONT
    run.font.size = size
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), EN_FONT)
    rf.set(qn("w:eastAsia"), ZH_FONT)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_doc(doc, lang):
    st = doc.styles["Normal"]
    st.font.name = EN_FONT
    st.font.size = SIZE
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), EN_FONT)
    rf.set(qn("w:eastAsia"), ZH_FONT)
    pf = st.paragraph_format
    pf.line_spacing = 1.15
    pf.space_after = Pt(3)
    pf.space_before = Pt(0)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        h = doc.styles[name]
        h.font.name = EN_FONT
        h.font.size = SIZE
        h.font.bold = True
        h.font.italic = False
        h.font.color.rgb = None
        hr = h.element.get_or_add_rPr()
        hf = hr.find(qn("w:rFonts"))
        if hf is None:
            hf = OxmlElement("w:rFonts")
            hr.insert(0, hf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            hf.set(qn(a), ZH_FONT if a == "w:eastAsia" else EN_FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if hf.get(qn(a)) is not None:
                del hf.attrib[qn(a)]
        h.paragraph_format.space_before = Pt(9 if name == "Heading 1" else 6)
        h.paragraph_format.space_after = Pt(3)
        h.paragraph_format.keep_with_next = True
        col = hr.find(qn("w:color"))
        if col is not None:
            hr.remove(col)
    for s in doc.sections:
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.left_margin = s.right_margin = Cm(2.54)
        s.top_margin = s.bottom_margin = Cm(2.54)


def para(doc, text, lang, indent=True, align=None, bold=False, italic=False, size=SIZE, space_after=None, keep=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(24) if lang == "zh" else Cm(1.27)
    if align:
        p.alignment = align
    elif lang == "zh":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if keep:
        p.paragraph_format.keep_with_next = True
    add_rich(p, text, bold, italic, size)
    return p


def add_rich(p, text, bold=False, italic=False, size=SIZE):
    """支援 *斜體* 與 **粗體** 標記（APA 英文統計符號、期刊名斜體用）。"""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            set_run_font(p.add_run(part[2:-2]), size, True, italic)
        elif part.startswith("*"):
            set_run_font(p.add_run(part[1:-1]), size, bold, True)
        else:
            set_run_font(p.add_run(part), size, bold, italic)


def heading(doc, text, level):
    h = doc.add_heading(level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 and text.startswith(("摘要", "Abstract")) else WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(h.add_run(text), SIZE, True, False)
    return h


def page_number_footer(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run)
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), tag)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = txt
        run._r.append(e)
    pg = section._sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        section._sectPr.append(pg)
    pg.set(qn("w:start"), "1")


def figure(doc, n, file, title, note, lang, width=15.0):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    set_run_font(p.add_run(f"圖 {n}" if lang == "zh" else f"Figure {n}"), SIZE, True)
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(title), SIZE, False, lang == "en")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(os.path.join(FIG, f"{file}_{lang}.png"), width=Cm(width))
    if note:
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(8)
        add_rich(q, ("註：" if lang == "zh" else "*Note.* ") + note)


def _border(cell, edges):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for edge, val in edges.items():
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), val)
        e.set(qn("w:sz"), "8")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        b.append(e)


def table(doc, n, title, header, rows, note, lang, widths=None, size=SIZE, label=None):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    lab = label or n
    set_run_font(p.add_run(f"表 {lab}" if lang == "zh" else f"Table {lab}"), SIZE, True)
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(title), SIZE, False, lang == "en")
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = widths is None
    if widths:  # LibreOffice／Word 都讀 tblGrid，固定欄寬
        grid = t._tbl.tblGrid
        for gc, wcm in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(wcm * 567)))
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        t._tbl.tblPr.append(lay)
    for j, h in enumerate(header):
        c = t.cell(0, j)
        c.text = ""
        add_rich(c.paragraphs[0], h, size=size)
        _border(c, {"top": "single", "bottom": "single"})
    for i, r in enumerate(rows, start=1):
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.text = ""
            add_rich(c.paragraphs[0], str(v), size=size)
            if i == len(rows):
                _border(c, {"bottom": "single"})
    for i, row in enumerate(t.rows):
        for j, c in enumerate(row.cells):
            for pp in c.paragraphs:
                pp.paragraph_format.first_line_indent = Pt(0)
                pp.paragraph_format.space_after = Pt(0)
                pp.paragraph_format.line_spacing = 1.0
                txt = c.text.strip()
                numeric = bool(re.fullmatch(r"[\d.,–—()（）%=C ≈<>I²無None-]*", txt)) and txt != ""
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER if (j > 0 and (numeric or i == 0)) else WD_ALIGN_PARAGRAPH.LEFT
            if widths:
                c.width = Cm(widths[j])
    if note:
        q = doc.add_paragraph()
        q.paragraph_format.space_before = Pt(3)
        q.paragraph_format.space_after = Pt(8)
        add_rich(q, ("註：" if lang == "zh" else "*Note.* ") + note)
    return t


def references(doc, N, lang):
    for _, r in N["refs"].iterrows():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.27)
        p.paragraph_format.first_line_indent = Cm(-1.27)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        ref = r.apa
        if lang == "en" and ref.startswith("MD. Piece 研究團隊"):
            ref = (ref.replace("MD. Piece 研究團隊", "MD. Piece Research Team")
                   .replace("慢性病縱向復發預測：MD. Piece 復發引擎 × 真實縱向資料與數位雙生 [未出版研究資料]",
                            "Longitudinal prediction of chronic disease recurrence: MD. Piece recurrence engine with real longitudinal data and a digital twin [Unpublished raw data]")
                   .replace("md_piece 虛擬患者模擬器與 16 種疾病設定 [電腦軟體]", "md_piece virtual patient simulator and 16 disease settings [Computer software]")
                   .replace("GitHub 儲存庫", "GitHub repository").replace("，分支", ", branch"))
        add_rich(p, _italic_journal(ref))


def _italic_journal(ref):
    """APA：期刊名與卷號斜體。"""
    m = re.search(r"\)\. (.+?)\. ([^.]+?), (\d+[^,]*?)(\(|,|\.)", ref)
    if not m or "GitHub" in ref:
        return ref
    j, vol = m.group(2), re.match(r"\d+", m.group(3)).group(0)
    return ref.replace(f"{j}, {vol}", f"*{j}, {vol}*", 1)


def cover(doc, lang, title, subtitle):
    zh = lang == "zh"
    for _ in range(4):
        doc.add_paragraph()
    para(doc, "中華民國中小學科學展覽會" if zh else "Science Fair", lang, False, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "作品說明書" if zh else "Project Report", lang, False, WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_after=36)
    para(doc, "作品名稱" if zh else "Project Title", lang, False, WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    para(doc, title, lang, False, WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    para(doc, subtitle, lang, False, WD_ALIGN_PARAGRAPH.CENTER, space_after=48)
    for k in (["科　　別：", "組　　別：", "關 鍵 詞：慢性病復發、縱向研究、文獻統整、虛擬患者、預測模型", "編　　號："] if zh else
              ["Category:", "Division:", "Keywords: chronic disease recurrence, longitudinal study, literature synthesis, virtual patients, prediction model", "Entry number:"]):
        para(doc, k, lang, False, WD_ALIGN_PARAGRAPH.LEFT)


# ───────────── 組裝 ─────────────
def build(lang):
    import report_text
    N = numbers()
    blocks = report_text.blocks(N, lang)
    doc = Document()
    style_doc(doc, lang)
    cover(doc, lang, *report_text.TITLE[lang])
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page_number_footer(sec)
    doc.sections[0].footer.is_linked_to_previous = False
    figs = tabs = 0
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            heading(doc, b[1], 1)
        elif kind == "h2":
            heading(doc, b[1], 2)
        elif kind == "p":
            para(doc, b[1], lang)
        elif kind == "pn":
            para(doc, b[1], lang, indent=False)
        elif kind == "fig":
            figs += 1
            n = int(re.match(r"fig(\d+)", b[1]).group(1))  # 圖號與檔名一致，內文引用寫的就是這個號碼
            assert n == figs, f"圖 {n} 出現的順序不對（第 {figs} 張）"
            figure(doc, n, b[1], b[2], b[3], lang, b[4] if len(b) > 4 else 15.0)
        elif kind == "tab":
            label = b[7] if len(b) > 7 else None
            tabs += 0 if label else 1
            table(doc, tabs, b[1], b[2], b[3], b[4], lang, b[5] if len(b) > 5 else None, b[6] if len(b) > 6 else SIZE, label)
        elif kind == "refs":
            references(doc, N, lang)
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    name = report_text.FILENAME[lang]
    path = os.path.join(DOCS, name)
    doc.save(path)
    return path


if __name__ == "__main__":
    for lang in ("zh", "en"):
        print(build(lang))
