# -*- coding: utf-8 -*-
"""獨立文件「腎損傷九種病因與機轉文獻回顧」——數字全由結果檔、書目全由 PubMed 快照產生，不手打。
    python build_nine_causes_review.py [--force]
輸入：results/nine_causes_nhanes.json、params/nine_causes_plan.json、docs/nine_causes_refs.json（選用理由與中文摘要）、
      docs/nine_causes_pubmed.json（PubMed 書目快照）。
輸出：Obsidian 主稿（研究計畫書/腎損傷九種病因與機轉文獻回顧.md）、repo 副本（experiments/docs/）、
      文獻總表（研究修訂_20260926/腎損傷九種病因文獻總表.xlsx）。Word：python -X utf8 md2docx.py <主稿> --fair"""
import json
import os
import re
import shutil
import subprocess
import sys

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT)
import build_fair_v3_2 as F            # noqa: E402  APA 文中引用（render_citations）與路徑
from build_plan_v3 import write        # noqa: E402  已在 Obsidian 修改過的主稿預設不覆蓋

NAME = "腎損傷九種病因與機轉文獻回顧"
XLSX = r"C:\Users\tpc10\Desktop\研究修訂_20260926\腎損傷九種病因文獻總表.xlsx"
FIG = "圖一_九種病因調整盛行率比.png"                                   # 有圖才放（figures/nine_causes/）
J = lambda *p: json.load(open(os.path.join(ROOT, *p), encoding="utf-8"))
R, PLAN = J("results", "nine_causes_nhanes.json"), J("params", "nine_causes_plan.json")
REFS, PM = J("docs", "nine_causes_refs.json"), J("docs", "nine_causes_pubmed.json")["articles"]
P, PH = R["primary"], R["phenotype"]
HIV_DONE = "status" not in P["HIV"]
git = lambda *a: subprocess.run(["git", *a], cwd=ROOT, capture_output=True, text=True, check=True).stdout.strip()
PLAN_COMMIT = git("log", "--format=%h", "--diff-filter=A", "--", "params/nine_causes_plan.json").splitlines()[-1]
RESULT_COMMIT = git("log", "-1", "--format=%h", "--", "results/nine_causes_nhanes.json")

# ── APA 書目（由 PubMed 快照排版）
JOURNAL = {
    "Advances in chronic kidney disease": "Advances in Chronic Kidney Disease",
    "American journal of kidney diseases : the official journal of the National Kidney Foundation": "American Journal of Kidney Diseases",
    "American journal of nephrology": "American Journal of Nephrology",
    "American journal of physiology. Renal physiology": "American Journal of Physiology-Renal Physiology",
    "Annals of hepatology": "Annals of Hepatology", "Annals of internal medicine": "Annals of Internal Medicine",
    "Annals of the rheumatic diseases": "Annals of the Rheumatic Diseases", "Archives of internal medicine": "Archives of Internal Medicine",
    "Arthritis & rheumatology (Hoboken, N.J.)": "Arthritis & Rheumatology", "Arthritis and rheumatism": "Arthritis & Rheumatism",
    "Arthritis research & therapy": "Arthritis Research & Therapy", "BMC nephrology": "BMC Nephrology", "BMJ (Clinical research ed.)": "BMJ",
    "Clinical infectious diseases : an official publication of the Infectious Diseases Society of America": "Clinical Infectious Diseases",
    "Clinical journal of the American Society of Nephrology : CJASN": "Clinical Journal of the American Society of Nephrology",
    "Diabetes care": "Diabetes Care", "Digestive diseases and sciences": "Digestive Diseases and Sciences",
    "Hepatology (Baltimore, Md.)": "Hepatology", "Hypertension (Dallas, Tex. : 1979)": "Hypertension", "JAMA": "JAMA",
    "JAMA internal medicine": "JAMA Internal Medicine",
    "Journal of the American Society of Nephrology : JASN": "Journal of the American Society of Nephrology",
    "Kidney international": "Kidney International", "Lancet (London, England)": "The Lancet", "Nature genetics": "Nature Genetics",
    "Nature reviews. Disease primers": "Nature Reviews Disease Primers", "Nature reviews. Nephrology": "Nature Reviews Nephrology",
    "Nature reviews. Rheumatology": "Nature Reviews Rheumatology", "PLoS medicine": "PLOS Medicine", "PloS one": "PLOS ONE",
    "Rheumatology (Oxford, England)": "Rheumatology", "Seminars in arthritis and rheumatism": "Seminars in Arthritis and Rheumatism",
    "The Journal of infectious diseases": "The Journal of Infectious Diseases",
    "The New England journal of medicine": "The New England Journal of Medicine"}
KEEP = ["United Kingdom Prospective Diabetes Study", "National Kidney Foundation", "International Society of Nephrology",
        "Renal Pathology Society", "National Institutes of Health", "European League Against Rheumatism",
        "American College of Rheumatology", "International Chapel Hill Consensus Conference", "Oxford Classification",
        "IgA Nephropathy Classification Working Group", "Kidney Disease: Improving Global Outcomes", "Controversies Conference",
        "HIV Medicine Association", "Infectious Diseases Society of America", "National Health and Nutrition Examination Survey",
        "United States", "Bayesian", "Consensus Conference", "D:A:D"]
TITLE_FIX = {"38182286": "KDIGO 2024 clinical practice guideline for the management of lupus nephritis.",   # PubMed 標題含全大寫字
             "36410841": "KDIGO 2022 clinical practice guideline for the prevention, diagnosis, evaluation, and treatment of "
                         "hepatitis C in chronic kidney disease."}


def sentence_case(pmid, t):
    """APA 文章標題：句首與冒號後大寫，其餘小寫；縮寫、內含大寫字母或數字之詞、專有名詞片語、單一字母 B/C 保留。"""
    if pmid in TITLE_FIX:
        return TITLE_FIX[pmid]
    t = " ".join(t.split()).rstrip(".")
    prot = [m.span() for ph in KEEP for m in re.finditer(r"(?<![A-Za-z])" + re.escape(ph) + r"(?![A-Za-z])", t)]
    out, first = [], True
    for m in re.finditer(r"[A-Za-z][A-Za-z0-9']*|[^A-Za-z]+", t):
        w, (i, j) = m.group(0), m.span()
        if not w[0].isalpha():
            out.append(w)
            if re.search(r":\s*$", w):
                first = True
            elif re.search(r"\d", w):
                first = False
            continue
        if any(a <= i and j <= b for a, b in prot) or re.search(r"[A-Z0-9]", w[1:]) or w in ("B", "C") or (w.isupper() and len(w) > 1):
            out.append(w)
        else:
            out.append(w[0].upper() + w[1:] if first else w.lower())
        first = False
    return "".join(out) + ("" if t.endswith("?") else ".")


def page_range(p):
    """PubMed 縮寫頁碼 → APA 完整頁碼（602-10 → 602–610；e96-138 → e96–e138）。"""
    if not p or "-" not in p:
        return p
    a, b = p.split("-", 1)
    ma, mb = re.match(r"^([A-Za-z]*)(\d+)$", a), re.match(r"^([A-Za-z]*)(\d+)$", b)
    if not (ma and mb):
        return f"{a}–{b}"
    pre, na = ma.groups()
    nb = mb.group(2)
    nb = na[:len(na) - len(nb)] + nb if len(nb) < len(na) else nb
    return f"{pre}{na}–{pre}{nb}"


def apa(r):
    a = PM[r["pmid"]]
    if r.get("group_author"):
        authors = r["group_author"] + "."
    else:
        names = [f"{last}, " + " ".join(f"{c}." for c in ini) for last, ini in a["authors"]]
        if len(names) == 1:
            authors = names[0]
        elif len(names) <= 20:
            authors = ", ".join(names[:-1]) + ", & " + names[-1]
        else:
            authors = ", ".join(names[:19]) + ", … " + names[-1]
    vol = f"*{JOURNAL[a['journal']]}, {a['volume']}*" + (f"({a['issue']})" if a.get("issue") else "")
    return f"{authors} ({r['year']}). {sentence_case(r['pmid'], a['title'])} {vol}, {page_range(a['pages'])}. https://doi.org/{a['doi']}"


def in_text(r):
    if r.get("group_author"):
        return r["au"]
    last = [x[0] for x in PM[r["pmid"]]["authors"]]
    return last[0] if len(last) == 1 else f"{last[0]} & {last[1]}" if len(last) == 2 else f"{last[0]} et al."


REF = {}
for r in REFS["refs"]:
    a = PM[r["pmid"]]
    assert a["journal"] in JOURNAL, a["journal"]
    assert not any("Retract" in t for t in a["article_types"]), r["key"]
    sort = r["group_author"] if r.get("group_author") else " ".join(f"{l} {i}" for l, i in a["authors"])
    REF[r["key"]] = dict(au=in_text(r), date=str(r["year"]), sort=sort, apa=apa(r))
    if r.get("group_author"):
        REF[r["key"]]["long"] = r["group_author"].replace(" (KDIGO) ", " [KDIGO] ")
F.REF.clear()
F.REF.update(REF)

# ── 數字（全由結果檔）
DISP = {"糖尿病": "糖尿病", "肥胖": "肥胖", "高尿酸血症": "高尿酸血症", "痛風": "痛風", "B型肝炎": "B 型肝炎", "C型肝炎": "C 型肝炎", "HIV": "HIV"}
pct = lambda p: f"{100 * p:.1f}%" if p >= 0.02 else f"{100 * p:.2f}%"
tab = lambda c: f"{c['est']:.2f}（{c['ci95'][0]:.2f}–{c['ci95'][1]:.2f}）"          # 表格（欄名已註明 95% CI）
rr = lambda c: f"{c['est']:.2f}（95% CI {c['ci95'][0]:.2f}–{c['ci95'][1]:.2f}）"
rri = lambda c: f"{c['est']:.2f}，95% CI {c['ci95'][0]:.2f}–{c['ci95'][1]:.2f}"   # 用在括號內，避免括號套括號
g = lambda d, k="腎損傷", W=P: W[d]["groups"][k]
apr = lambda d, k="腎損傷", W=P: rr(W[d]["contrasts"][k]["aPR"])
apri = lambda d, k="腎損傷", W=P: rri(W[d]["contrasts"][k]["aPR"])
aprt = lambda d, k="腎損傷", W=P: tab(W[d]["contrasts"][k]["aPR"])
crude = lambda d: f"{P[d]['contrasts']['腎損傷']['PR_crude']['est']:.2f}"
cnt = lambda n: f"{n:,}"
DONE = [d for d in P if "status" not in P[d]]
EXPECT = {"糖尿病": "腎損傷者較常見", "肥胖": "腎損傷者較常見", "高尿酸血症": "腎損傷者較常見", "痛風": "腎損傷者較常見",
          "B型肝炎": "無法區分", "C型肝炎": "腎損傷者較常見", "HIV": "腎損傷者較常見"}
for d, v in EXPECT.items():   # 與計畫事前寫下之文獻預期逐字對照
    assert PLAN["expectations_from_literature"][d].startswith("aPR>1" if v == "腎損傷者較常見" else "不確定"), d
MATCH = {d: P[d]["verdict"] == EXPECT[d] for d in DONE}
assert all(MATCH.values()), MATCH          # 若有不一致，下面的文字敘述須改寫，不得照用
PH_N = {k: PH["糖尿病"]["groups"][k]["n"] for k in ("只有白蛋白尿", "只有eGFR<60", "兩者皆有")}
C = R["cohort"]
n_type = lambda *ts: sum(1 for r in REFS["refs"] if r["type"] in ts)
N_REF = len(REFS["refs"])


def results_table():
    rows = ["| 疾病 | 腎損傷者：陽性／人數；加權盛行率（95% CI） | 無腎損傷者：陽性／人數；加權盛行率（95% CI） | 粗盛行率比 | 調整盛行率比（95% CI） | 判讀 | 文獻預期 |",
            "|---|---|---|---|---|---|---|"]
    for d in P:
        if d not in DONE:
            rows.append(f"| {DISP[d]} | 檢驗檔待下載 | — | — | — | — | 較常見 |")
            continue
        cell = lambda k: (f"{cnt(g(d, k)['n_events'])}／{cnt(g(d, k)['n'])}；{pct(g(d, k)['prevalence'])}"
                          f"（{pct(g(d, k)['ci95'][0])}–{pct(g(d, k)['ci95'][1])}）")
        rows.append(f"| {DISP[d]} | {cell('腎損傷')} | {cell('無腎損傷')} | {crude(d)} | {aprt(d)} | {P[d]['verdict']} | "
                    f"{'較常見' if EXPECT[d] == '腎損傷者較常見' else '無法區分'}（{'符合' if MATCH[d] else '不符'}） |")
    return "\n".join(rows)


def phenotype_table():
    rows = ["| 疾病 | 只有白蛋白尿 | 只有 eGFR<60 | 兩者皆有 |", "|---|---|---|---|"]
    for d in DONE:
        rows.append(f"| {DISP[d]} | " + " | ".join(aprt(d, k, PH) for k in ("只有白蛋白尿", "只有eGFR<60", "兩者皆有")) + " |")
    return "\n".join(rows)


hiv_sentence = (f"腎損傷成人 HIV 陽性 {pct(g('HIV')['prevalence'])}，無腎損傷者 {pct(g('HIV', '無腎損傷')['prevalence'])}，"
                f"調整盛行率比 {apr('HIV')}，判讀為「{P['HIV']['verdict']}」。" if HIV_DONE else
                "NHANES 1999–2018 有 HIV 檢驗（只驗 18–49 歲，2009 年起 18–59 歲），檢驗檔待下載後依同一計畫補算。")
hiv_limit = (f"HIV 只驗 20–49 歲（2009 年起 20–59 歲），腎損傷者中陽性僅 {cnt(g('HIV')['n_events'])} 人。" if HIV_DONE
             else "HIV 待檢驗檔下載後補算。")
hiv_concl = (f"HIV 亦是腎損傷者較常見（調整盛行率比 {apri('HIV')}），以白蛋白尿型態為主（只有白蛋白尿 {apri('HIV', '只有白蛋白尿', PH)}；兩者皆有 {apri('HIV', '兩者皆有', PH)}）。" if HIV_DONE
             else "HIV 待補。")
abstract_hiv = f"HIV 亦較常見（{apri('HIV')}），但陽性人數少、區間較寬。" if HIV_DONE else "HIV 之公開資料檢驗待檢驗檔下載後補上。"
FIG_MD = (f"\n![[figure/nine_causes/{FIG}|620]]\n\n**圖一、腎損傷成人之九種病因調整盛行率比**（NHANES 1999–2018；"
          "左：腎損傷對無腎損傷；右：依腎臟型態對兩項皆正常者；橫線為 95% CI）\n"
          if os.path.exists(os.path.join(ROOT, "figures", "nine_causes", FIG)) else "")

BODY = f"""# 已知腎損傷，可能是哪一種病？

## ——九種病因與機轉之文獻回顧及公開資料檢驗

## 摘要

腎損傷（eGFR<60 mL/min/1.73 m² 或尿白蛋白／肌酸酐比 ACR ≥30 mg/g）只告訴我們腎臟受傷，沒有告訴我們原因。本文回顧代謝、免疫、感染三個方向各三種病——糖尿病、肥胖、高尿酸血症／痛風；紅斑性狼瘡、IgA 腎病變、ANCA 相關血管炎；B 型肝炎、C 型肝炎、HIV——說明它們為什麼會傷腎、會留下什麼腎臟線索、已知腎損傷時要驗什麼才找得到，共引用 {N_REF} 篇經 PubMed 核對之文獻（臨床指引 {n_type('臨床指引')} 篇、系統性回顧與統合分析 {n_type('系統性回顧／統合分析')} 篇）。再以美國國家健康與營養調查（NHANES）1999–2018 年公開資料、事前提交分析計畫，檢驗其中量得到的病：調整年齡、性別與種族後，腎損傷成人之糖尿病（調整盛行率比 {apri('糖尿病')}）、痛風（{apri('痛風')}）、高尿酸血症（{apri('高尿酸血症')}）、C 型肝炎（{apri('C型肝炎')}）與肥胖（{apri('肥胖')}）都比沒有腎損傷的成人常見，B 型肝炎無法區分（{apri('B型肝炎')}），{len(DONE)} 項皆與事前寫下之文獻預期一致。{abstract_hiv}腎損傷的型態也是線索：高尿酸與痛風集中在 eGFR 下降者，糖尿病與 C 型肝炎在白蛋白尿者較明顯。免疫方向三種病在 NHANES 量不到，須靠專項抗體檢驗與腎臟切片。

## 壹、前言

### 一、問題：知道腎受傷，還要知道為什麼

慢性腎臟病的評估除了 GFR 與白蛋白尿分級，還要判定病因[@kdigo_ckd2024]，因為處置取決於病因：糖尿病要控制血糖並使用保護腎臟的藥物[@kdigo_dm2022]，狼瘡性腎炎與 ANCA 相關血管炎需要免疫抑制治療[@anders2020,hellmich2024]，病毒相關腎病則以抗病毒治療為主[@gupta2015]。本文從「已知腎損傷」反過來問：可能是哪一個方向的哪一種病？為什麼？

### 二、為什麼選這九種病

每個方向選三種。代謝方向選一般人口最常見、與腎損傷關係最密切的糖尿病、肥胖與高尿酸血症／痛風；免疫方向選紅斑性狼瘡、IgA 腎病變與 ANCA 相關血管炎，三者都會造成腎絲球腎炎、常需腎臟切片確診，其中 IgA 腎病變是最常見的原發性腎絲球腎炎[@lai2016]，ANCA 相關血管炎是急進性腎絲球腎炎最常見的原因[@berden2010]；感染方向選會長期感染、可用血液檢驗確認的 B 型肝炎、C 型肝炎與 HIV。

### 三、方法

#### （一）文獻回顧

以 PubMed 檢索，優先選用國際臨床指引、系統性回顧與統合分析、大型世代研究與隨機對照試驗，輔以病理分類與機轉研究，每種病 7–9 篇（部分指引與綜述跨病共用）。每篇先以引文比對取得 PMID，再逐篇取回 PubMed 書目資料，核對標題、作者、期刊卷頁與 DOI（2026 年 10 月 9 日）；文末參考文獻由該 PubMed 書目快照以程式排版，不手打。核對時發現 1 篇已被撤稿（ANCA 相關血管炎之 avacopan 試驗，PubMed 標記為撤稿），予以剔除；另有 6 篇因與其他文獻重複或以治療為主而不採用。逐篇之類型、主要發現與 PMID／DOI 見文獻總表（Excel）。

#### （二）公開資料檢驗

取 NHANES 1999–2018 年十個週期之 20 歲以上成人，腎臟狀態可判定者 {cnt(C['kidney_known'])} 人：腎損傷（eGFR<60 或 ACR≥30）{cnt(C['kidney_damage'])} 人、無腎損傷 {cnt(C['no_kidney_damage'])} 人；eGFR 以不含種族係數之 CKD-EPI 2021 公式計算[@inker2021]。疾病定義見表二。比較兩組之加權盛行率，並以加權邏輯斯迴歸調整年齡（十歲一組）、性別與種族，用邊際標準化求調整盛行率比；95% 信賴區間以刪一 PSU 摺刀法（{C['psu']} 組複製權重）估計，盛行率採 Korn–Graubard 區間。判讀規則事前寫定：調整盛行率比之 95% 信賴區間下界大於 1 為「腎損傷者較常見」，上界小於 1 為「較少見」，其餘為「無法區分」。分析計畫、程式與各病之文獻預期於計算前提交（提交代碼 {PLAN_COMMIT}），結果另行提交（{RESULT_COMMIT}）。

**表二、NHANES 之疾病定義**

| 疾病 | 定義 | 週期 |
|---|---|---|
| 糖尿病 | 曾被醫師告知有糖尿病，或 HbA1c ≥6.5% | 1999–2018 |
| 肥胖 | BMI ≥30 kg/m²（檢查時懷孕者不判定） | 1999–2018 |
| 高尿酸血症 | 血清尿酸男性 >7.0、女性 >5.7 mg/dL[@zhu2011] | 1999–2018 |
| 痛風 | 曾被醫師告知有痛風 | 2007–2018 |
| B 型肝炎 | B 型肝炎表面抗原（HBsAg）陽性 | 1999–2018 |
| C 型肝炎 | HCV RNA 陽性（抗體陰性者未驗 RNA，視為陰性） | 1999–2018 |
| HIV | HIV 抗體陽性（2015 年起為抗原／抗體篩檢加確認檢驗）；只驗 18–49 歲（2009 年起 18–59 歲） | 1999–2018 |
| 紅斑性狼瘡、IgA 腎病變、ANCA 相關血管炎 | NHANES 沒有診斷問卷、專項抗體與腎臟切片，無法判定 | — |

## 貳、九種病總覽

**表一、九種病的傷腎方式與腎臟線索**

| 方向 | 疾病 | 為什麼會傷腎 | 典型腎臟病變 | 已知腎損傷時的線索 | 確認方法 |
|---|---|---|---|---|---|
| 代謝 | 糖尿病 | 高血糖造成高過濾、腎絲球肥大與硬化 | 糖尿病腎病變（結節性腎絲球硬化） | 白蛋白尿逐年增加後 eGFR 下降；也可只有 eGFR 下降 | HbA1c、空腹血糖 |
| 代謝 | 肥胖 | 高過濾、腎絲球內壓上升、腎內脂質堆積 | 肥胖相關腎絲球病變（腎絲球肥大，可合併局部節段性硬化） | 蛋白尿為主，少有腎病症候群，進展較慢 | BMI；先排除其他病因 |
| 代謝 | 高尿酸血症／痛風 | 動物研究：血管病變與血壓上升；腎功能下降也使尿酸排泄減少（互為因果） | 無特定腎絲球病變；痛風者常合併腎結石 | 以 eGFR 下降為主 | 血清尿酸、痛風病史 |
| 免疫 | 紅斑性狼瘡 | 抗核抗體與免疫複合體沉積、補體活化 | 狼瘡性腎炎（I–VI 級） | 蛋白尿、血尿，多在診斷後 5 年內；常有全身表現 | 抗核抗體、抗 dsDNA、補體；腎臟切片 |
| 免疫 | IgA 腎病變 | 異常 IgA1 免疫複合體沉積於系膜 | 系膜增生（牛津 MEST-C 評分） | 顯微或肉眼血尿合併蛋白尿；多為青壯年 | 腎臟切片（唯一確診方法） |
| 免疫 | ANCA 相關血管炎 | ANCA 活化嗜中性球，造成壞死性小血管炎 | 壞死性新月體腎絲球腎炎 | 數週至數月內 eGFR 快速下降、血尿；多為年長者；常合併呼吸道症狀 | ANCA（MPO、PR3）；腎臟切片 |
| 感染 | B 型肝炎 | 病毒抗原免疫複合體沉積 | 膜性腎病變為主 | 蛋白尿（可達腎病症候群）、顯微血尿 | HBsAg、病毒量 |
| 感染 | C 型肝炎 | 冷凝球蛋白免疫複合體沉積 | 膜增生性腎絲球腎炎（冷凝球蛋白血症） | 蛋白尿、血尿、高血壓、低補體 | 抗體、HCV RNA；冷凝球蛋白 |
| 感染 | HIV | 病毒直接感染腎臟上皮；APOL1 基因；抗病毒藥物 | HIV 相關腎病變（塌陷型腎絲球病變）等 | 大量蛋白尿、腎功能快速惡化；或用藥後 eGFR 緩降 | HIV 抗原／抗體檢驗 |

## 參、各病：為什麼會傷腎、會留下什麼線索

### 一、代謝方向

#### （一）糖尿病

1. 為什麼會傷腎：長期高血糖使腎絲球處於高過濾狀態，腎絲球肥大、系膜擴張，最後腎絲球硬化，並伴隨腎小管間質發炎與纖維化[@alicic2017]；腎臟是糖尿病微血管損傷最重要的標的[@thomas2015]。病理上依序可見腎絲球基底膜增厚、系膜擴張、結節性硬化（Kimmelstiel-Wilson 結節）與廣泛腎絲球硬化[@tervaert2010]。
2. 會留下什麼線索：典型病程是高過濾、白蛋白尿逐年增加，之後 eGFR 下降、走向末期腎臟病[@alicic2017]；第 2 型糖尿病診斷後每年約 2% 出現微白蛋白尿，10 年時約四分之一[@adler2003]。但美國資料顯示近年糖尿病成人白蛋白尿比例下降、eGFR 下降比例上升[@afkarian2016]，所以沒有白蛋白尿不能排除糖尿病的影響。
3. 有多常見：約 40% 糖尿病患者會發生糖尿病腎病，是全球慢性腎臟病的首要原因[@alicic2017]，在已開發國家約占末期腎臟病的一半[@tuttle2014]；美國 2009–2014 年約 26% 糖尿病成人有糖尿病腎病[@afkarian2016]。腎損傷一旦出現，eGFR 與白蛋白尿對死亡及末期腎臟病的相對風險，在有無糖尿病者相近[@fox2012]。
4. 已知腎損傷時：驗 HbA1c 或空腹血糖即可找到，確認後依指引處置[@kdigo_dm2022]。
5. 公開資料：腎損傷成人有糖尿病者 {pct(g('糖尿病')['prevalence'])}，無腎損傷者 {pct(g('糖尿病', '無腎損傷')['prevalence'])}，調整盛行率比 {apr('糖尿病')}。依腎臟型態，「兩者皆有」（{apri('糖尿病', '兩者皆有', PH)}）與「只有白蛋白尿」（{apri('糖尿病', '只有白蛋白尿', PH)}）最高，「只有 eGFR<60」較低（{apri('糖尿病', '只有eGFR<60', PH)}），與糖尿病腎病多經白蛋白尿表現一致（表四）。

#### （二）肥胖

1. 為什麼會傷腎：肥胖使腎絲球過濾率、腎血漿流量、過濾分率與腎小管鈉再吸收都上升[@dagati2016]；重度肥胖者腎絲球過濾率高出 51%、腎血漿流量高出 31%，入球小動脈擴張使動脈壓直接傳到腎絲球微血管[@chagnac2000]。腎絲球為了應付高過濾而肥大，形成適應性局部節段性腎絲球硬化；脂肪組織分泌的激素與腎內異位脂質堆積，也讓足細胞產生胰島素阻抗[@dagati2016]。
2. 會留下什麼線索：肥胖相關腎絲球病變以腎絲球肥大為特徵，與原發性局部節段性腎絲球硬化相比，腎病症候群較少（5.6% 對 54%）、病程較緩[@kambham2001]；多數為穩定或緩慢進展的蛋白尿，但最多三分之一會走向腎衰竭，而減重能減少蛋白尿並逆轉高過濾[@dagati2016]。
3. 有多常見、證據多強：腎切片中肥胖相關腎絲球病變的比例 15 年間增加十倍（0.2% 到 2.0%）[@kambham2001]。32 萬名成人追蹤，BMI 越高末期腎臟病風險越高，BMI≥40 者為正常體重者的 7.07 倍，調整血壓與糖尿病後仍存在[@hsu2006]；119 萬名 17 歲青少年追蹤約 25 年，肥胖者末期腎臟病風險為 6.89 倍，非糖尿病性者亦為 3.41 倍[@vivante2012]。統合分析中肥胖者腎病風險為 1.83 倍[@wang2008]；546 萬人之個人資料統合分析中，BMI 35 對 25 之腎功能下降風險為 1.69 倍，再調整其他共病後降為 1.28 倍[@chang2019]，表示部分影響經由糖尿病、高血壓等共病。
4. 已知腎損傷時：量 BMI。肥胖常與糖尿病、高血壓同時存在，要認定是肥胖本身造成，須先排除其他病因，必要時以腎臟切片確認[@kambham2001,dagati2016]。
5. 公開資料：腎損傷成人肥胖 {pct(g('肥胖')['prevalence'])}，無腎損傷者 {pct(g('肥胖', '無腎損傷')['prevalence'])}，調整盛行率比 {apr('肥胖')}，是量得到的病中最小的，且三種腎臟型態相近（表四），符合「常見但影響中等」的危險因子。腎病晚期可能體重下降，會讓這個關聯被低估。

#### （三）高尿酸血症／痛風

1. 為什麼會傷腎，以及為什麼腎損傷會讓尿酸升高：尿酸與腎臟互為因果。一方面，腎功能些微變化就會改變血尿酸，高尿酸常被視為腎功能下降的標記[@johnson2018,kang2002]；另一方面，動物研究顯示輕度高尿酸會讓血壓上升、腎內出現缺血型損傷，而且不需要尿酸結晶[@mazzali2001]；在已受損的腎臟中，高尿酸會加重蛋白尿、腎絲球硬化與間質纖維化，機轉與 COX-2 及血栓素造成的血管病變有關，降尿酸可以預防[@kang2002]。
2. 會留下什麼線索：痛風盛行率隨腎功能惡化而上升：無慢性腎臟病者 2–3%、第 3 期 11–13%、第 4 期超過 30%，調整尿酸後仍存在[@juraschek2013]；痛風患者中 24% 有第 3 期以上慢性腎臟病、14% 有腎結石[@roughley2015]。
3. 證據多強：世代研究之統合分析顯示高尿酸是新發慢性腎臟病的獨立預測因子（勝算比 2.35）[@li2014]，痛風與慢性腎臟病之調整勝算比為 2.41[@roughley2015]。但兩個隨機對照試驗以 allopurinol 降尿酸，在第 3–4 期慢性腎臟病[@badve2020]與第 1 型糖尿病腎病[@doria2020]都沒有減緩腎功能下降；孟德爾隨機化研究也多不支持尿酸本身造成腎病[@johnson2018]。因此在已知腎損傷的人身上，高尿酸較可能是「腎功能下降的結果，可能再加重損傷」，而不是主要病因。
4. 已知腎損傷時：驗血清尿酸、詢問痛風病史。高尿酸在腎損傷者很常見，但不宜據此認定為病因，仍應尋找其他原因。
5. 公開資料：腎損傷成人高尿酸血症 {pct(g('高尿酸血症')['prevalence'])}，無腎損傷者 {pct(g('高尿酸血症', '無腎損傷')['prevalence'])}，調整盛行率比 {apr('高尿酸血症')}；痛風 {pct(g('痛風')['prevalence'])} 對 {pct(g('痛風', '無腎損傷')['prevalence'])}，調整盛行率比 {apr('痛風')}。依腎臟型態，高尿酸血症在「只有 eGFR<60」（{apri('高尿酸血症', '只有eGFR<60', PH)}）與「兩者皆有」（{apri('高尿酸血症', '兩者皆有', PH)}）遠高於「只有白蛋白尿」（{apri('高尿酸血症', '只有白蛋白尿', PH)}），痛風亦同（表四）——集中在過濾功能下降者，正是「腎功能下降使尿酸排泄減少」這個反向機轉的痕跡。

### 二、免疫方向

#### （一）紅斑性狼瘡

1. 為什麼會傷腎：紅斑性狼瘡是對自身細胞核抗原失去免疫耐受的自體免疫病，產生抗核抗體；死亡細胞釋出的核酸像病毒一樣，經 Toll 樣受體活化第一型干擾素；抗體結合腎內抗原並活化補體，形成免疫複合體腎絲球腎炎[@lech2013]。
2. 會留下什麼線索：狼瘡性腎炎依腎臟切片分為 I–VI 級：系膜型（I、II）、局部或瀰漫增生型（III、IV）、膜性（V）與硬化型（VI）[@weening2004]，2018 年修訂再加入活動度與慢性度指數[@bajema2018]。多數發生在紅斑性狼瘡診斷後 5 年內，而且常是第一個表現[@anders2020]。
3. 有多常見、多嚴重：國際起始世代 1,827 名患者中，38.3% 發生狼瘡性腎炎，八成在收案時就已存在；腎炎者 10 年內 10.1% 進展到末期腎臟病，死亡風險為 2.98 倍[@hanly2016]。統合分析顯示，已開發國家 5 年末期腎臟病風險由 1970 年代的 16% 降到 1990 年代中期約 11% 後持平，第 IV 級 15 年風險高達 44%[@tektonidou2016]。
4. 已知腎損傷時：若同時有發燒、皮膚黏膜、關節或血球異常等全身表現，應驗抗核抗體（紅斑性狼瘡分類之必要條件），再驗抗 dsDNA 等特異抗體與補體[@aringer2019]；確診與分級需腎臟切片，治療依指引[@kdigo_ln2024]。
5. 公開資料：NHANES 1999–2018 沒有紅斑性狼瘡之診斷問卷，也沒有抗 dsDNA、補體或腎臟切片，無法檢驗。主研究曾以 1999–2004 年剩餘血清之抗核抗體作為免疫方向之代理標籤，但抗核抗體陽性不等於紅斑性狼瘡，本文不重做。

#### （二）IgA 腎病變

1. 為什麼會傷腎：以「多重打擊」解釋：體內產生鉸鏈區半乳糖缺乏之 IgA1（部分由基因決定）；身體再產生針對它的自體抗體；兩者形成免疫複合體並沉積於腎絲球系膜；系膜細胞因而增生、分泌細胞激素與基質，造成腎損傷[@suzuki2011]，並經由補體與細胞間訊號傷及足細胞與腎小管間質[@lai2016]。全基因組研究找到主要組織相容複合體與補體因子 H 相關基因（CFHR1/CFHR3）等易感區域，風險基因頻率與亞洲、歐洲、非洲族群之盛行率差異平行[@gharavi2011]。
2. 會留下什麼線索：表現從無症狀的顯微血尿到肉眼血尿，常合併蛋白尿[@lai2016,wyatt2013]；病理以牛津分類 MEST-C 評分（系膜增生、內皮增生、節段硬化、腎小管萎縮／間質纖維化、新月體）[@trimarchi2017]。
3. 有多常見、多嚴重：全球最常見的原發性腎絲球腎炎；30–40% 患者在發病 20–30 年後進展到末期腎臟病[@lai2016]。英國登錄 2,439 名患者之腎臟存活中位數 11.4 年，即使蛋白尿低於 0.44 g/g，10 年內仍約兩成腎衰竭[@pitcher2023]；結合 eGFR、血壓、蛋白尿與切片評分之國際預測工具，C 統計量 0.82[@barbour2019]。
4. 已知腎損傷時：青壯年出現血尿合併蛋白尿，應想到 IgA 腎病變；血液檢驗無法確診，必須腎臟切片[@kdigo_gd2021]。
5. 公開資料：NHANES 沒有腎臟切片，無法檢驗。

#### （三）ANCA 相關血管炎

1. 為什麼會傷腎：身體對嗜中性球的兩種蛋白——蛋白酶 3（PR3）或髓過氧化酶（MPO）——產生自體抗體（ANCA）[@kitching2020]。ANCA 活化已被致敏的嗜中性球與單核球，啟動補體替代路徑，引發呼吸爆發、脫顆粒與嗜中性球胞外網，造成小血管壞死性發炎[@jennette2014]。依國際命名共識，它屬於幾乎沒有免疫沉積之小血管炎，包括顯微多血管炎、肉芽腫性多血管炎與嗜酸性肉芽腫性多血管炎[@jennette2013]。
2. 會留下什麼線索：腎臟出現壞死性新月體腎絲球腎炎；ANCA 相關血管炎是全球急進性腎絲球腎炎最常見的原因，切片可分局部型、新月體型、混合型與硬化型，並預測腎臟預後[@berden2010]。上、下呼吸道與腎臟最常受侵犯[@kitching2020]。
3. 有多嚴重：倫敦 246 名腎臟受侵犯之患者，中位年齡 66 歲，診斷時肌酸酐中位數 3.87 mg/dL，92% ANCA 陽性，28% 進展到末期腎衰竭[@booth2003]；發生率隨年齡增加[@watts2022]。
4. 已知腎損傷時：年長者腎功能在數週到數月內快速惡化、合併血尿或呼吸道症狀時，應驗 ANCA 並安排腎臟切片[@hellmich2024,kdigo_aav2024]。
5. 公開資料：NHANES 沒有 ANCA 檢驗，無法檢驗。

### 三、感染方向

#### （一）B 型肝炎

1. 為什麼會傷腎：病毒相關腎絲球腎炎是指活躍的病毒複製直接造成腎絲球病變，與鏈球菌感染之後才發生的腎炎不同[@kupin2017]；B 型肝炎病毒抗原與抗體形成免疫複合體沉積於腎絲球，腎組織中可測到病毒抗原[@bhimma2004]，也可能引起結節性多動脈炎[@johnson1990]。
2. 會留下什麼線索：最常見膜性腎病變，表現為不同程度之蛋白尿與顯微血尿[@gupta2015]，可達腎病症候群[@bhimma2004]。兒童常在 e 抗原清除後緩解[@bhimma2004]；成人則少自發緩解，香港 21 名成人患者約三分之一緩慢惡化，29% 腎功能變差、10% 需長期透析[@lai1991]。
3. 證據多強：全球 B 型肝炎帶原者逾 4 億人[@kupin2017]。統合分析中，縱貫研究顯示 B 型肝炎感染者末期腎臟病風險為 3.87 倍，但橫斷面調查未見 B 型肝炎與慢性腎臟病或蛋白尿相關[@fabrizi2017]。
4. 已知腎損傷時：出現膜性腎病變或腎病症候群者應驗 B 型肝炎表面抗原（HBsAg），陽性再驗病毒量；相關腎病以抗病毒治療為主[@gupta2015,terrault2018,kdigo_gd2021]。
5. 公開資料：腎損傷成人 HBsAg 陽性 {pct(g('B型肝炎')['prevalence'])}，無腎損傷者 {pct(g('B型肝炎', '無腎損傷')['prevalence'])}，調整盛行率比 {apr('B型肝炎')}，無法區分，與統合分析中橫斷面研究之結果一致[@fabrizi2017]。美國 HBsAg 陽性者少（{cnt(P['B型肝炎']['n_events'])} 人），檢定力有限；在 B 型肝炎流行地區[@lai1991]結果可能不同。

#### （二）C 型肝炎

1. 為什麼會傷腎：C 型肝炎病毒刺激 B 細胞增生，產生具類風濕因子活性之 IgM，與病毒及抗病毒 IgG 結合成冷凝球蛋白（低於體溫會沉澱的免疫球蛋白），沉積於腎絲球[@johnson1993,roccatello2018]；混合型冷凝球蛋白血症多數由 C 型肝炎引起[@roccatello2018]。
2. 會留下什麼線索：最常見第 1 型膜增生性腎絲球腎炎，多合併第 2 型混合型冷凝球蛋白血症，表現為蛋白尿、顯微血尿、高血壓、腎炎或腎病症候群[@gupta2015]，常有低補體[@johnson1993]。
3. 證據多強：統合分析中，C 型肝炎使新發慢性腎臟病風險為 1.43 倍、蛋白尿 1.51 倍[@fabrizi2015]；美國退伍軍人世代中，70 歲以下 C 型肝炎抗體陽性者末期腎臟病風險為 2.80 倍[@tsui2007]，另一個 102 萬人世代中腎功能下降為 1.15 倍、末期腎臟病 1.98 倍[@molnar2015]。直接抗病毒藥物清除病毒後，冷凝球蛋白腎炎患者之肌酸酐與蛋白尿改善[@sise2016]，支持病毒是原因；但 C 型肝炎也是慢性腎臟病的併發症（透析族群較易感染），關係是雙向的[@perico2009]。
4. 已知腎損傷時：慢性腎臟病評估時應篩檢 C 型肝炎，先驗抗體，陽性再驗 HCV RNA；膜增生性病變者另驗冷凝球蛋白與補體[@kdigo_hcv2022,gupta2015]。
5. 公開資料：腎損傷成人 HCV RNA 陽性 {pct(g('C型肝炎')['prevalence'])}，無腎損傷者 {pct(g('C型肝炎', '無腎損傷')['prevalence'])}；未調整時兩組相近（粗盛行率比 {crude('C型肝炎')}，信賴區間含 1），調整年齡、性別與種族後為 {apr('C型肝炎')}，表示這些因素在兩組分布不同而遮蔽了關聯。依腎臟型態，「只有白蛋白尿」（{apri('C型肝炎', '只有白蛋白尿', PH)}）與「兩者皆有」（{apri('C型肝炎', '兩者皆有', PH)}）較高，「只有 eGFR<60」則無法區分（{apri('C型肝炎', '只有eGFR<60', PH)}），與 C 型肝炎以腎絲球病變（蛋白尿）為主之傷腎方式一致。

#### （三）HIV

1. 為什麼會傷腎：HIV 會直接感染腎絲球與腎小管之上皮細胞，腎臟成為病毒的儲存處[@bruggeman2000,rosenberg2015]；帶兩個 APOL1 風險基因（只見於非洲裔）者發生 HIV 相關腎病變之勝算比為 29，這類人若感染 HIV 而未治療，發生風險約 50%[@kopp2011]。此外，抗病毒藥物 tenofovir 及 ritonavir 加強之 atazanavir、lopinavir 累積使用，與腎功能下降有關[@ryom2013]。
2. 會留下什麼線索：HIV 感染者之腎病包括 HIV 相關腎病變（典型為塌陷型腎絲球病變）、非塌陷型局部節段性腎絲球硬化、免疫複合體腎病，以及藥物與共病造成之腎病[@swanepoel2018,rosenberg2015]；抗病毒治療普及後，長期治療者之腎硬化與糖尿病腎病變增加[@rosenberg2015]。
3. 有多常見：全球 HIV 感染者之慢性腎臟病（eGFR<60）盛行率為 4.8%–6.4%，非洲最高[@ekrikpo2018]；歐洲世代之慢性腎臟病發生率為每千人年 6.2，傳統與 HIV 相關危險因子都有預測力[@mocroft2015]。
4. 已知腎損傷時：HIV 感染者應定期檢查腎功能與尿蛋白[@lucas2014]；反過來，原因不明、尤其合併大量蛋白尿之腎損傷，HIV 是需要排除的病因之一[@swanepoel2018]。
5. 公開資料：{hiv_sentence}

## 肆、公開資料檢驗結果

### 一、腎損傷成人中，哪些病比較常見

**表三、腎損傷與無腎損傷成人之疾病盛行率（NHANES 1999–2018）**

{results_table()}

註：陽性／人數為未加權人數；盛行率經抽樣權重加權，95% CI 為 Korn–Graubard 區間；調整變項為年齡（十歲一組）、性別與種族；「判讀」依事前規則；「文獻預期」為計算前寫入分析計畫之預期方向。
{FIG_MD}
在已檢驗的 {len(DONE)} 項中，腎損傷成人較常見的依調整盛行率比排序為：{'、'.join(f'{DISP[d]}（{P[d]["contrasts"]["腎損傷"]["aPR"]["est"]:.2f}）' for d in sorted(DONE, key=lambda d: -P[d]['contrasts']['腎損傷']['aPR']['est']) if P[d]['verdict'] == '腎損傷者較常見')}；B 型肝炎無法區分。以盛行率而言，腎損傷成人中最常見的是肥胖（{pct(g('肥胖')['prevalence'])}）、高尿酸血症（{pct(g('高尿酸血症')['prevalence'])}）與糖尿病（{pct(g('糖尿病')['prevalence'])}），病毒性肝炎都不到 2%。全部結果與事前寫下之文獻預期方向一致。

### 二、腎損傷的型態提供線索

**表四、依腎臟型態之調整盛行率比（對照組：eGFR 與 ACR 皆正常）**

{phenotype_table()}

註：數值為調整盛行率比（95% CI），調整變項同表三；只有 eGFR 或 ACR 其中一項有值者不列入。各型態人數以糖尿病分析為例：只有白蛋白尿 {cnt(PH_N['只有白蛋白尿'])} 人、只有 eGFR<60 {cnt(PH_N['只有eGFR<60'])} 人、兩者皆有 {cnt(PH_N['兩者皆有'])} 人。

兩種型態指向不同的病：高尿酸血症與痛風集中在 eGFR 下降者，與「過濾減少使尿酸排泄減少」一致，也提醒尿酸較可能是腎損傷的結果；糖尿病與 C 型肝炎在白蛋白尿者較明顯，與兩者以腎絲球病變傷腎一致；肥胖在三種型態相近。

### 三、限制

1. 橫斷面資料只能說「同時存在」，不能說誰造成誰；尿酸與肥胖還受腎功能反向影響。
2. 腎損傷依單次檢驗判定，未經 3 個月確認[@kdigo_ckd2024]，可能納入暫時性異常。
3. 盛行率低的病（B、C 型肝炎與 HIV）陽性人數少，信賴區間較寬；{hiv_limit}
4. 免疫方向三種病無法以 NHANES 量測，只能依文獻說明。
5. 美國資料：各病盛行率因地而異，例如 B 型肝炎在流行地區遠較常見[@kupin2017,lai1991]，外推需謹慎。

## 伍、結論：已知腎損傷，可以怎麼想

1. **代謝方向最常見**：糖尿病是最常見也最確定的病因，腎損傷成人約 {round(100 * g('糖尿病')['prevalence'])}% 有糖尿病，調整後為無腎損傷者的 {P['糖尿病']['contrasts']['腎損傷']['aPR']['est']:.1f} 倍；肥胖很常見但影響中等；高尿酸與痛風多半是腎功能下降的結果，而降尿酸並未減緩腎病[@badve2020,doria2020]。
2. **感染方向找到了就能治**：C 型肝炎與腎損傷相關，且以白蛋白尿型態為主；B 型肝炎在美國資料中看不出差異；{hiv_concl}病毒相關腎病以抗病毒治療為主[@gupta2015]，C 型肝炎病毒清除後腎臟可以改善[@sise2016]。
3. **免疫方向少見但急**：紅斑性狼瘡、IgA 腎病變與 ANCA 相關血管炎可在數年（甚至數週）內導致末期腎臟病[@hanly2016,pitcher2023,booth2003]，須靠專項抗體與腎臟切片確診，一般常規檢驗無法確診。
4. **看腎損傷的型態找線索**：以白蛋白尿或蛋白尿為主，想到糖尿病、肥胖、病毒相關腎絲球病變與免疫性腎炎；以 eGFR 下降為主而尿酸偏高，尿酸多為結果，仍須找其他病因；腎功能短期內快速惡化合併血尿，想到 ANCA 相關血管炎與狼瘡性腎炎，應儘快接受專科評估與切片[@kdigo_aav2024,kdigo_ln2024]。

## 陸、參考文獻資料

{{REFERENCES}}
"""


def emit(fmt):
    text, refs = F.render_citations(BODY)
    text = text.replace("{REFERENCES}", "\n\n".join(refs))
    if fmt == "repo":
        text = text.replace(f"![[figure/nine_causes/{FIG}|620]]", f"![](../kidney_cause/figures/nine_causes/{FIG})")
    assert "{" not in text and "[@" not in text, "仍有未處理之欄位或引用"
    assert len(refs) == N_REF, (len(refs), N_REF)          # 每篇文獻都有被引用
    return text, refs


def excel():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "文獻總表"
    head = ["編號", "方向", "疾病", "文中引用", "年份", "標題（PubMed）", "期刊", "卷（期）：頁", "文獻類型", "主要發現（中文摘要）",
            "PMID", "DOI", "PubMed 連結", "核對"]
    ws.append(head)
    group = {"糖尿病": "代謝", "肥胖": "代謝", "高尿酸血症／痛風": "代謝", "紅斑性狼瘡": "免疫", "IgA 腎病變": "免疫",
             "ANCA 相關血管炎": "免疫", "B 型肝炎": "感染", "C 型肝炎": "感染", "HIV": "感染", "總論": "總論"}
    for i, r in enumerate(REFS["refs"], 1):
        a = PM[r["pmid"]]
        ref = REF[r["key"]]
        ws.append([i, "、".join(dict.fromkeys(group[d] for d in r["diseases"])), "、".join(r["diseases"]),
                   f"{ref['au']}（{ref['date']}）", r["year"], " ".join(a["title"].split()), JOURNAL[a["journal"]],
                   f"{a['volume']}" + (f"({a['issue']})" if a.get("issue") else "") + f"：{page_range(a['pages'])}",
                   r["type"], r["finding"], int(r["pmid"]), a["doi"], f"https://pubmed.ncbi.nlm.nih.gov/{r['pmid']}/",
                   "PubMed 書目已核對 2026-10-09"])
    widths = [6, 8, 16, 28, 7, 60, 28, 16, 16, 70, 10, 30, 40, 22]
    for col, w in zip("ABCDEFGHIJKLMN", widths):
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    s = wb.create_sheet("各病篇數")
    s.append(["疾病", "篇數", "其中臨床指引", "其中系統性回顧／統合分析"])
    for d in group:
        rs = [r for r in REFS["refs"] if d in r["diseases"]]
        s.append([d, len(rs), sum(r["type"] == "臨床指引" for r in rs), sum(r["type"] == "系統性回顧／統合分析" for r in rs)])
    s.append(["不重複總計", N_REF, n_type("臨床指引"), n_type("系統性回顧／統合分析")])
    x = wb.create_sheet("剔除")
    x.append(["PMID", "理由"])
    for e in REFS["_meta"]["excluded"]:
        x.append([int(e["pmid"]), e["why"]])
    for sh in (s, x):
        sh.column_dimensions["A"].width, sh.column_dimensions["B"].width = 22, 90
        for c in sh[1]:
            c.font = Font(bold=True)
    wb.save(XLSX)
    print(f"[寫出] {XLSX}（{N_REF} 篇）")


def main():
    vault, _ = emit("vault")
    repo, _ = emit("repo")
    fm = ("---\ntitle: 腎損傷九種病因與機轉文獻回顧\ndate: 2026-10-09\ntags: [腎臟研究, 文獻回顧, 科展]\n待辦:\n"
          + ("" if HIV_DONE else "  - HIV 檢驗檔（10 檔約 0.96 MB）待使用者同意下載後補算\n")
          + "  - 經使用者確認後，再併入說明書與計畫書\n---\n\n")
    note = ("> [!info] 由 `experiments/kidney_cause/build_nine_causes_review.py` 自結果檔與 PubMed 書目快照產生（數字與書目不手打）。"
            "可直接在此筆記修改（重跑產生程式預設不覆蓋已修改的筆記）；Word 以 `md2docx.py <本檔> --fair` 轉出。\n\n")
    src = os.path.join(ROOT, "figures", "nine_causes", FIG)
    if os.path.exists(src):                                   # Obsidian 主稿與 Word 轉檔都從 vault 的 figure/ 讀圖
        os.makedirs(os.path.join(F.VAULT, "figure", "nine_causes"), exist_ok=True)
        shutil.copy2(src, os.path.join(F.VAULT, "figure", "nine_causes", FIG))
    write(os.path.join(F.VAULT, f"{NAME}.md"), fm + note + vault, guard=True)
    write(os.path.join(F.REPO_DOCS, f"{NAME}.md"), repo, guard=False)
    excel()
    print(f"文獻 {N_REF} 篇｜已檢驗疾病 {len(DONE)}｜HIV {'已' if HIV_DONE else '未'}計算｜計畫 {PLAN_COMMIT}｜結果 {RESULT_COMMIT}")


if __name__ == "__main__":
    main()
