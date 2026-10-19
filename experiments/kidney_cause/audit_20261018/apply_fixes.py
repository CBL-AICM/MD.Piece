# -*- coding: utf-8 -*-
"""依 2026-10-18 全面審核報告修改兩份文件，另存新檔（原檔不動）。
每處以原文定位：找不到或找到多處即停。能在單一 run 內換就不動格式；跨 run 才整段重排（保留第一個 run 的格式）。
用法：python apply_fixes.py <B原檔> <B新檔> <A原檔> <A新檔> <圖一.png> <圖七.png>"""
import copy
import re
import sys

import docx
from docx.oxml.ns import qn
from PIL import Image

B_SRC, B_DST, A_SRC, A_DST, FIG1, FIG7 = sys.argv[1:7]
for s, d in [(B_SRC, B_DST), (A_SRC, A_DST)]:
    assert s != d, "不覆寫原檔"


def all_paras(d):
    yield from d.paragraphs
    for t in d.tables:
        for row in t.rows:
            seen = set()
            for c in row.cells:
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                yield from c.paragraphs


def rep(d, old, new, log, label):
    hits = [p for p in all_paras(d) if old in p.text]
    # 合併儲存格會讓同一段落出現多次，以底層元素去重
    uniq = list({id(p._p): p for p in hits}.values())
    assert len(uniq) == 1, f"[{label}] 找到 {len(uniq)} 處：{old[:40]}"
    p = uniq[0]
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            log.append((label, old, new))
            return p
    full = p.text.replace(old, new, 1)
    p.runs[0].text = full
    for r in p.runs[1:]:
        r.text = ""
    log.append((label, old, new))
    return p


def split_caption(d, glued, caption, template_text, log, label):
    """把黏在段尾的圖說拆成獨立段落，格式複製自 template_text 那段圖說。"""
    p = rep(d, glued + caption, glued, log, label)
    tmpl = next(q for q in d.paragraphs if q.text.strip() == template_text)
    newp = copy.deepcopy(tmpl._p)
    p._p.addnext(newp)
    np_ = docx.text.paragraph.Paragraph(newp, p._parent)
    np_.runs[0].text = caption
    for r in np_.runs[1:]:
        r.text = ""


def caption_before_table(d, table_first_cell, caption, template_text, log, label):
    t = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == table_first_cell)
    tmpl = next(q for q in d.paragraphs if q.text.strip() == template_text)
    newp = copy.deepcopy(tmpl._p)
    t._tbl.addprevious(newp)
    np_ = docx.text.paragraph.Paragraph(newp, d._body)
    np_.runs[0].text = caption
    for r in np_.runs[1:]:
        r.text = ""
    log.append((label, "（無圖說）", caption))


def swap_image(d, caption_prefix, png, log, label):
    """換掉圖說 caption_prefix 之前最近一張圖；寬不變，高依新圖比例（圖一要求比例不變）。"""
    body = list(d.element.body.iterchildren())
    idx = next(i for i, el in enumerate(body) if el.tag == qn("w:p") and
               "".join(t.text or "" for t in el.iter(qn("w:t"))).strip().startswith(caption_prefix))
    for el in reversed(body[:idx]):
        blips = el.findall(".//" + qn("a:blip"))
        if blips:
            break
    rid = blips[0].get(qn("r:embed"))
    part = d.part.related_parts[rid]
    old_w, old_h = Image.open(__import__("io").BytesIO(part.blob)).size
    new_w, new_h = Image.open(png).size
    part._blob = open(png, "rb").read()
    ext = el.find(".//" + qn("wp:extent"))
    if abs(old_h / old_w - new_h / new_w) > 1e-3:
        cx = int(ext.get("cx"))
        cy = int(cx * new_h / new_w)
        assert cy <= int(ext.get("cy")) * 1.001, f"[{label}] 新圖比原圖高，會改變分頁"
        ext.set("cy", str(cy))
        for e2 in el.iter(qn("a:ext")):
            e2.set("cy", str(cy))
    log.append((label, f"圖檔 {old_w}×{old_h}", f"新圖 {new_w}×{new_h}"))


def zh_len(s):
    return len(re.sub(r"\s", "", s))


# ======================= 說明書 =======================
B = docx.Document(B_SRC)
LB = []
rep(B, "the seven measurable in NHANES all differed in prevalence in the directions expected from the literature",
    "the seven measurable in NHANES all matched the directions expected from the literature", LB, "1 英文摘要")
rep(B, "檢驗評估陽性／人數", "事後評估陽性／人數", LB, "2 表十六")
# 表十六第一區塊列名補回「事前指定」，與「事後評估」成對（只換該表第 2 列第 1 欄）
t16 = next(t for t in B.tables if any("事後評估陽性／人數" in c.text for r in t.rows for c in r.cells))
c = t16.rows[1].cells[0]
assert c.text.strip() == "陽性／人數", c.text
c.paragraphs[0].runs[0].text = "事前指定評估陽性／人數"
for r in c.paragraphs[0].runs[1:]:
    r.text = ""
LB.append(("2 表十六", "陽性／人數", "事前指定評估陽性／人數"))
swap_image(B, "圖七、", FIG7, LB, "2 圖七")
swap_image(B, "圖一、", FIG1, LB, "5 圖一")
rep(B, "並把結果做成工具（目的五至七），最後匯整為結論。右側的連線表示目的六直接使用目的一的資料基礎，不經過預測模型。目的四中的切片轉錄體比較與目的八（九種候選病因）是在上述分析完成後加入的補充，未畫入圖一。",
    "把結果做成工具，並比較九種候選病因（目的五至八），最後匯整為結論。左側的連線表示目的六、右側的連線表示目的八，兩者都直接使用目的一的資料基礎，不經過預測模型。目的四中的切片轉錄體比較（資料來自 NIH GEO）與目的八，是在 NHANES 分析完成後加入的補充。",
    LB, "5 圖一說明")
# 6 中文摘要 ≦300 字
ab_old = next(p.text for p in B.paragraphs if p.text.startswith("腎損傷的病因可分為感染、免疫、代謝三個方向"))
edits = [("腎損傷的病因可分為感染、免疫、代謝三個方向。", "腎損傷病因分感染、免疫、代謝三方向。"),
         ("檢驗健檢常規抽血、驗尿能否回推病因方向", "檢驗常規抽血、驗尿能否回推病因方向"),
         ("（病毒性肝炎）0.807，訊號主要來自 C 型肝炎；", "（病毒性肝炎）0.807，主要來自 C 型肝炎；"),
         ("另以 NIH 公開的病理確診腎臟切片轉錄體，在同一處理批次內比較，", "另以 NIH 公開之病理確診切片轉錄體做同批次事後比較，"),
         ("九種候選病因中可量測的七項，在腎損傷者的盛行率皆與文獻預期一致。", "九種候選病因中可測的七項，盛行率比皆符合文獻預期。"),
         ("免疫方向的訊號在腎臟組織，", "免疫方向訊號在腎臟組織，")]
for o, n in edits:
    rep(B, o, n, LB, "6 中文摘要")
ab_new = next(p.text for p in B.paragraphs if p.text.startswith("腎損傷病因分感染"))
assert zh_len(ab_new) <= 300, zh_len(ab_new)
LB.append(("6 中文摘要", f"{zh_len(ab_old)} 字", f"{zh_len(ab_new)} 字"))
rep(B, "2017 年公開", "2018 年公開", LB, "8 表二")
rep(B, "約 433 MB", "約 432 MB", LB, "9 容量")
split_caption(B, "以自行設計下列機制（表七）。", "表七、研究品質管控機制", "表六、各週期檢驗標定之處理", LB, "10 表七圖說")
rep(B, "以自行設計下列機制（表七）。", "本研究設計下列機制（表七）。", LB, "10 表七前句文法")
caption_before_table(B, "階段", "表二十一、研究階段與進度", "表六、各週期檢驗標定之處理", LB, "10 表二十一圖說")
rep(B, "本作品已界定常規檢驗回推病因方向的能力與邊界；", "本作品已界定常規檢驗回推病因方向的能力與邊界（表二十一）；", LB, "10 表二十一引用")
rep(B, "本研究只使用公開、已去識別的檔案，不接觸任何受試者。",
    "NHANES 調查協定由 NCHS 倫理審查委員會核准（NCHS, 2026）；本研究只使用公開、已去識別的檔案，不接觸任何受試者。", LB, "11 倫理引用")
rep(B, "ERCB 的糖尿病腎病變是需要切片的非典型個案；IgA 腎病變與高血壓腎病變無法評估。",
    "ERCB 的糖尿病腎病變是需要切片的非典型個案；IgA 腎病變與高血壓腎病變無法評估。另有一例（GSM2810968）樣本名稱標為膜性腎病，診斷欄卻為紅斑性狼瘡，本研究依診斷欄歸類，只影響腎小管間質重現的一例。",
    LB, "14 限制")
B.save(B_DST)

# ======================= 研究報告 =======================
A = docx.Document(A_SRC)
LA = []
swap_image(A, "圖 7", FIG7, LA, "2 圖 7")
rep(A, "判定干擾素訊號為狼瘡特有。", "判定干擾素路徑足以區分狼瘡腎炎與膜性腎病變。", LA, "3 方法")
rep(A, "腎小管間質為 0.975，符合計算前寫定的「干擾素訊號為狼瘡特有」判定。兩者同屬免疫複合體沉積的腎病，差異主要在干擾素訊號。",
    "腎小管間質為 0.975，符合計算前寫定的判定條件，表示干擾素路徑可區分狼瘡腎炎與膜性腎病變；但單核巨噬細胞（0.942）與補體（0.903）也幾乎一樣高，兩者同屬免疫複合體沉積的腎病，差異不只在干擾素。",
    LA, "3 結果")
rep(A, "干擾素路徑反而較低（0.41）。",
    "干擾素路徑反而較低（0.41）；以上單路徑數字未經交叉驗證與多重比較校正，只作描述。", LA, "4 結果")
P194_OLD = next(p.text for p in A.paragraphs if p.text.startswith("這組結果在臨床上有兩點意義。"))
P194_NEW = ("這組結果有兩點值得留意。第一，相對於 FSGS 與微小變化病，糖尿病腎病變切片的補體、單核巨噬細胞與 T 細胞路徑分數也較高，"
            "與糖尿病腎病變有先天免疫與補體活化的文獻相符（Flyvbjerg, 2017; Tesch, 2017），這可能是 ANCA 相關血管炎對糖尿病腎病變時"
            "免疫路徑模型無法區分的原因之一。不過這是每批只有 5 至 7 例、未經交叉驗證與多重比較校正的描述，而且 Tesch（2017）指出糖尿病腎病變"
            "的免疫反應與自體免疫腎病明顯不同，因此「以發炎路徑區分免疫與代謝方向可能不夠」只能列為待驗證的假說。第二，結果最一致的是狼瘡腎炎的"
            "干擾素升高：不論對 FSGS 或膜性腎病變，在腎絲球與腎小管間質都在 0.95 以上；糖尿病腎病變的干擾素在腎絲球偏低（0.41、0.40），"
            "在腎小管間質卻沒有同樣的方向（0.51、0.68），不能當作一致的訊號。臨床上可先以常規檢驗與病史判斷代謝與感染方向，若出現血尿、"
            "蛋白尿快速增加或腎功能短期惡化，再以補體 C3、C4、抗雙股 DNA 抗體、ANCA 與腎臟切片確認免疫方向。血液中的干擾素相關指標能否反映"
            "腎臟的干擾素訊號，需要另以資料驗證。")
rep(A, P194_OLD, P194_NEW, LA, "4 討論")
rep(A, "免疫方向的訊號存在於腎臟組織的干擾素、補體與發炎細胞路徑，常規檢驗沒有反映這些路徑的項目；",
    "免疫方向的訊號存在於腎臟組織，其中最一致的是狼瘡腎炎的干擾素升高，補體與發炎細胞路徑在糖尿病腎病變也可見升高；常規檢驗沒有反映這些路徑的項目；",
    LA, "13 結論")
# 7（＋15）摘要：標明模型別與事後分析，並維持 ≦350 字（含符號）
ab_old = next(p.text for p in A.paragraphs if p.text.startswith("本研究探討常規血液與尿液檢驗能否回推腎損傷"))
rep(A, "改以NIH公開的病理確診切片轉錄體同批次比較，狼瘡腎炎與ANCA血管炎可與非免疫腎病區分（0.94與0.93）。",
    "另以NIH公開之病理確診切片轉錄體做同批次事後比較，狼瘡腎炎與ANCA血管炎可與非免疫腎病區分（路徑模型0.94、全基因模型0.93）。",
    LA, "7 摘要")
rep(A, "以美國 NHANES 1999至2018年55,081名成人，校正檢驗量尺後，於腎臟指標異常者以巢狀交叉驗證評估。",
    "以美國 NHANES 1999至2018年55,081名成人，校正檢驗量尺後，於腎臟指標異常者評估。", LA, "7 摘要精簡")
rep(A, "免疫方向的訊號在腎臟組織，仍需專項檢驗與切片，不能取代臨床診斷。",
    "免疫方向訊號在腎臟組織，仍需專項檢驗與切片，不能取代臨床診斷。", LA, "7 摘要精簡")
rep(A, "本研究探討常規血液與尿液檢驗能否回推腎損傷", "本研究探討常規血液、尿液檢驗能否回推腎損傷", LA, "7 摘要精簡")
ab_new = next(p.text for p in A.paragraphs if p.text.startswith("本研究探討常規血液、尿液檢驗能否回推腎損傷"))
assert zh_len(ab_new) <= 350, zh_len(ab_new)
LA.append(("7/15 摘要字數", f"{zh_len(ab_old)} 字", f"{zh_len(ab_new)} 字（上限約 350 字含符號）"))
rep(A, "每例診斷皆經腎臟病理確認，對照組為活體捐腎者與腫瘤腎切除檢體。",
    "每例診斷皆經腎臟病理確認；資料集另含活體捐腎者與腫瘤腎切除檢體，本研究未納入比較。", LA, "12 GEO 資料")
rep(A, "IgA 腎病變與高血壓腎病變無法評估，也沒有獨立的外部資料集。",
    "IgA 腎病變與高血壓腎病變無法評估，也沒有獨立的外部資料集。另有一例（GSM2810968）樣本名稱標為膜性腎病，診斷欄卻為紅斑性狼瘡，本研究依診斷欄歸類，只影響腎小管間質重現的一例。",
    LA, "14 限制")
rep(A, "Hepatitis C RNA (HCV-RNA) and hepatitis C genotype", "Hepatitis C: RNA (HCV-RNA) and hepatitis C genotype", LA, "18 文獻冒號")
rep(A, "從原始檔重算切片轉錄體修正一的各項比較", "從原始檔重算切片轉錄體修正一與修正二的各項比較", LA, "附 品質管控表")
A.save(A_DST)

for name, L in [("說明書", LB), ("研究報告", LA)]:
    print(f"===== {name}：{len(L)} 處")
    for lab, o, n in L:
        print(f"[{lab}] {o[:50]} → {n[:70]}")
