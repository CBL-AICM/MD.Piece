# -*- coding: utf-8 -*-
"""研究數據原始名單：出處帳本（results/provenance.json）每一個原始檔＋GEO 切片逐例樣本名單 → Excel。
NHANES 檔之正式名稱取自 CDC 各檔說明頁（快取於 results/nhanes_file_titles.json）。
    python -X utf8 build_data_inventory.py <輸出.xlsx>"""
import gzip
import json
import os
import re
import sys
import time
import urllib.request

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
J = lambda *p: json.load(open(os.path.join(ROOT, *p), encoding="utf-8"))
PROV = J("results", "provenance.json")["files"]
MAN = J("params", "manifest.json")["files"]
CACHE = os.path.join(ROOT, "results", "nhanes_file_titles.json")
titles = json.load(open(CACHE, encoding="utf-8")) if os.path.exists(CACHE) else {}
OUT = sys.argv[1]

ROLE_ZH = {"demographics": "人口學（年齡、性別、種族、權重）"}


def cdc_title(url):
    doc = re.sub(r"\.xpt$", ".htm", url, flags=re.I)
    if doc in titles:
        return titles[doc]
    req = urllib.request.Request(doc, headers={"User-Agent": "Mozilla/5.0"})
    for _ in range(3):
        try:
            h = urllib.request.urlopen(req, timeout=30).read().decode("utf-8", "ignore")
            h3 = [re.sub(r"<[^>]+>", "", x).strip() for x in re.findall(r"<h3[^>]*>(.*?)</h3>", h, re.S)]
            t = next((x for x in h3 if x and x != "Table of Contents"), "")
            titles[doc] = re.sub(r"\s+", " ", t.replace("&amp;", "&"))
            time.sleep(0.3)
            return titles[doc]
        except Exception:
            time.sleep(2)
    titles[doc] = "（說明頁讀取失敗）"
    return titles[doc]


def cycle_of(url):
    m = re.search(r"/Public/(\d{4})/", url)
    if not m:
        return ""
    y = int(m.group(1))
    return "2021–2023" if y == 2021 else f"{y}–{y + 1}"


rows = []
for k, v in PROV.items():
    url = v["url"]
    if k.lower().endswith(".xpt"):
        cyc = cycle_of(url)
        group = "外部評估（2021–2023）" if cyc == "2021–2023" else "開發資料（1999–2018）"
        rows.append(["NHANES（CDC／NCHS）", cyc, k, cdc_title(url), MAN.get(k, {}).get("role", ""), group, url, v["bytes"], v["sha256"], v["downloaded_at_utc"]])
    else:
        if k.startswith("GSE104948"):
            t, cyc = "Glomerular transcriptome from European Renal cDNA Bank subjects and living donors", "腎絲球；" + k.split("-")[1].split(".")[0]
        elif k.startswith("GSE104954"):
            t, cyc = "Tubulointerstitial transcriptome from ERCB subjects with chronic kidney disease and living donor biopsies", "腎小管間質；" + k.split("-")[1].split(".")[0]
        else:
            t, cyc = "NCBI Gene：Homo sapiens 基因資訊（基因符號↔Entrez ID）", "—"
        rows.append(["NIH NCBI（GEO／Gene）", cyc, k, t, "immune_geo", "免疫方向補充（切片轉錄體）", url, v["bytes"], v["sha256"], v["downloaded_at_utc"]])
json.dump(titles, open(CACHE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
order = {"開發資料（1999–2018）": 0, "外部評估（2021–2023）": 1, "免疫方向補充（切片轉錄體）": 2}
rows.sort(key=lambda r: (order[r[5]], r[1], r[2]))

# ---- GEO 逐例樣本 ----
TASKS = {  # 同批次比較（修正一、修正二）：(批次, 陽性, 對照)
    "B1 ANCA 對非免疫": ("H7", {"ANCA"}, {"DN", "FSGS", "FSGS_MCD", "MCD"}),
    "B2 ANCA 對 DN": ("H7", {"ANCA"}, {"DN"}),
    "B3 SLE 對 FSGS": ("H5", {"SLE"}, {"FSGS"}),
    "C1 DN 對 FSGS／MCD": ("H7", {"DN"}, {"FSGS", "FSGS_MCD", "MCD"}),
    "C2 DN 對 MCD／TMD": ("H1", {"DN"}, {"MCD", "TMD"}),
    "C3 SLE 對 MN": ("H5", {"SLE"}, {"MN"}),
}
sys.path.insert(0, ROOT)
import immune_geo as ig  # noqa: E402  只用其診斷字串對應，結果與本名單無關

samples = []
for s, comp in [("GSE104948", "腎絲球"), ("GSE104954", "腎小管間質")]:
    for p in ["GPL22945", "GPL24120"]:
        path = os.path.join(ROOT, "data", "raw", "geo", f"{s}-{p}.txt.gz")
        meta = {}
        with gzip.open(path, "rt", encoding="utf-8") as f:
            for line in f:
                if line.startswith("!series_matrix_table_begin"):
                    break
                if line.startswith(("!Sample_geo_accession", "!Sample_title", "!Sample_source_name_ch1", "!Sample_characteristics_ch1")):
                    k, *v = line.rstrip("\n").split("\t")
                    v = [x.strip('"') for x in v]
                    if k == "!Sample_characteristics_ch1":
                        k = "tissue" if v[0].startswith("tissue") else "diagnosis"
                    meta[k] = v
        for i, gsm in enumerate(meta["!Sample_geo_accession"]):
            raw = meta["diagnosis"][i].replace("diagnosis:", "").strip()
            code = ig.dx_code(meta["diagnosis"][i])
            batch = re.search(r"(H\d+)-", meta["!Sample_source_name_ch1"][i]).group(1)
            used = [name for name, (bt, pos, neg) in TASKS.items() if batch == bt and code in pos | neg]
            role = "主分析" if comp == "腎絲球" else "重現"
            samples.append([gsm, s, comp, p, meta["!Sample_title"][i], batch, raw or "（空白＝活體捐腎者）", code,
                            "、".join(f"{u}（{role}）" for u in used) or "未納入同批次比較"])

wb = Workbook()
ws = wb.active
ws.title = "說明"
n_nh = sum(r[0].startswith("NHANES") for r in rows)
info = [
    ["研究數據原始名單"],
    ["資料來源", "全部為公開、已去識別之資料：美國 NHANES（CDC／NCHS）與 NIH NCBI（GEO 切片轉錄體、Gene 基因資訊）。"],
    ["檔案總數", f"{len(rows)} 個（NHANES {n_nh} 個：開發 {sum(r[5].startswith('開發') for r in rows)}、外部 {sum(r[5].startswith('外部') for r in rows)}；NIH NCBI {len(rows) - n_nh} 個）"],
    ["總容量", f"{sum(r[7] for r in rows) / 1e6:.1f} MB"],
    ["GEO 樣本", f"{len(samples)} 例（腎絲球 {sum(x[2] == '腎絲球' for x in samples)}、腎小管間質 {sum(x[2] == '腎小管間質' for x in samples)}）；同批次比較實際使用 {sum(x[8] != '未納入同批次比較' for x in samples)} 例"],
    ["如何驗證", "任何人可由「來源網址」重新下載，計算 SHA256 與本表比對；完全相同即為同一份檔案。研究程式只讀取帳本登錄且雜湊相符的檔案。"],
    ["來源紀錄", "experiments/kidney_cause/results/provenance.json（出處帳本）；本表由 build_data_inventory.py 產生。"],
    ["檔案正式名稱", "NHANES 取自 CDC 各檔說明頁之標題；GEO 取自系列標題。"],
]
for r in info:
    ws.append(r)
ws["A1"].font = Font(bold=True, size=14)
ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 110

for title, head, data, widths in [
    ("檔案名單", ["序號", "資料庫", "週期／系列", "檔名", "正式名稱", "用途代碼", "分組", "來源網址", "位元組", "SHA256", "下載時間（UTC）"],
     [[i + 1] + r for i, r in enumerate(rows)], [6, 20, 16, 22, 60, 16, 24, 70, 12, 66, 22]),
    ("GEO樣本名單", ["GSM 編號", "系列", "腎臟區域", "晶片平台", "樣本標題", "處理批次", "診斷（GEO 原文）", "診斷代碼", "納入之同批次比較"],
     samples, [12, 11, 11, 11, 22, 9, 52, 10, 60]),
]:
    w = wb.create_sheet(title)
    w.append(head)
    for c in w[1]:
        c.font = Font(bold=True)
    for r in data:
        w.append(r)
    for i, wd in enumerate(widths, 1):
        w.column_dimensions[get_column_letter(i)].width = wd
    w.freeze_panes = "A2"
    for row in w.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=False)
wb.save(OUT)
fails = [k for k, v in titles.items() if "失敗" in v]
print(f"ok：{len(rows)} 檔、{len(samples)} 例；CDC 標題讀取失敗 {len(fails)} 個")
