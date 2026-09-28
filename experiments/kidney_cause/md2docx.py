# -*- coding: utf-8 -*-
"""Obsidian 主稿（.md）→ 科展格式 Word，輸出於來源同目錄；圖以來源目錄為基準解析。
用法：python -X utf8 md2docx.py <主稿.md> [--fair]
只用已裝的 python-docx；語法覆蓋這份稿子實際用到的：#～#### 標題、段落、**粗體**、`code`、
[[wiki]]、[t](u)、> 引用、``` 圍欄、| 表格 |、1. 與 - 清單、![[figure/x.png]]（![[figure/x.png|480]]＝寬 480 px，以 1/96 吋計，與 Obsidian 同義）、---。
格式：A4、邊界 2.5 cm、內文 12 pt 標楷體＋Times New Roman、頁碼置中。
--fair：全國中小學科展作品說明書格式（附件六、七）——邊界 2 cm、新細明體、1.5 倍行高、內文 12 級、
主題（##）16 級粗體置中；「<!-- 分頁 -->」之前為封面（16 級、無頁碼），之後另起一節，頁碼自內文第一頁起算。
*斜體*（APA 之期刊名與卷號）亦支援。"""
import io, os, re, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

sys.stdout.reconfigure(encoding="utf-8")
FAIR = "--fair" in sys.argv
argv = [a for a in sys.argv[1:] if a != "--fair"]
if not argv:
    sys.exit("用法：python -X utf8 md2docx.py <主稿.md> [--fair]")
SRC = argv[0]
OUT = os.path.splitext(SRC)[0] + ".docx"
BASE = os.path.dirname(SRC)
CJK, LATIN, MONO, MONO_CJK = ("新細明體" if FAIR else "標楷體"), "Times New Roman", "Consolas", "細明體"
BODY_PT = 12
H_SIZE = {1: 16, 2: 16, 3: 14, 4: 12} if FAIR else {1: 20, 2: 16, 3: 14, 4: 12}

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.0 if FAIR else 2.5)


def rfonts(run, mono=False):
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    lat = MONO if mono else LATIN
    for k in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(k), lat)
    rf.set(qn("w:eastAsia"), MONO_CJK if mono else CJK)


def style_run(run, size=BODY_PT, bold=None, mono=False, color=None, italic=False):
    rfonts(run, mono)
    run.font.size = Pt(size)
    run.font.italic = italic or None
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


# Normal 樣式：12 pt、1.15 倍行距、段後 6 pt
normal = doc.styles["Normal"]
normal.font.size = Pt(BODY_PT)
normal.font.name = LATIN
normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
normal.paragraph_format.line_spacing = 1.5 if FAIR else 1.15
normal.paragraph_format.space_after = Pt(6)

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`|!?\[\[[^\]]+\]\]|\[[^\]]+\]\([^)]+\))")


def emit(par, text, size=BODY_PT, bold=False, mono=False):
    """把一段含行內標記的文字寫進段落。"""
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            emit(par, tok[2:-2], size, True, mono)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            style_run(par.add_run(tok[1:-1]), size, bold, mono, italic=True)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            style_run(par.add_run(tok[1:-1]), size, bold, True)
        elif tok.startswith("[[") or tok.startswith("![["):
            inner = tok[tok.index("[[") + 2:-2]
            style_run(par.add_run(inner.split("|")[-1]), size, bold, mono)
        elif tok.startswith("[") and "](" in tok:
            style_run(par.add_run(tok[1:tok.index("](")]), size, bold, mono)
        else:
            style_run(par.add_run(tok), size, bold, mono)


def shade(par, fill):
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def left_border(par):
    ppr = par._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    l = OxmlElement("w:left")
    for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "6"), ("w:color", "BFBFBF")):
        l.set(qn(k), v)
    b.append(l), ppr.append(b)


def cell_shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def heading(text, level, center=False):
    size = H_SIZE[level]
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt({1: 0, 2: 18, 3: 12, 4: 6}[level])
    p.paragraph_format.space_after = Pt(6)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    emit(p, text, size, True)
    return p


def paragraph(text, center=False, indent=False, size=BODY_PT, hanging=False):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif hanging:                                         # APA 參考文獻：靠左、凸排
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(0.85)
        p.paragraph_format.first_line_indent = Cm(-0.85)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.85)
    emit(p, text, size)
    return p


def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = text.startswith("**表")
    emit(p, text, 11)
    return p


def blockquote(lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    left_border(p)
    for i, l in enumerate(lines):
        if i:
            p.add_run().add_break()
        emit(p, l, 11)


def code_block(lines):
    for i, l in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Cm(0.3)
        shade(p, "F2F2F2")
        style_run(p.add_run(l if l else " "), 9, mono=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def list_item(text, marker):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.6)
    p.paragraph_format.space_after = Pt(3)
    style_run(p.add_run(marker + " "), BODY_PT)
    emit(p, text)


def table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c or "") for c in r)]
    ncol = max(len(r) for r in cells)
    fs = 10 if ncol >= 6 else 10.5
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblpr = t._tbl.tblPr                                  # 欄寬交給 Word 依內容自動調整，表寬＝版心寬
    tw = tblpr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW"); tblpr.append(tw)
    tw.set(qn("w:type"), "pct"); tw.set(qn("w:w"), "5000")
    for i, r in enumerate(cells):
        trpr = t.rows[i]._tr.get_or_add_trPr()            # 列不跨頁；跨頁時重複表頭
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            cell.paragraphs[0].paragraph_format.keep_with_next = i < len(cells) - 1 and len(cells) <= 15   # 短表不拆頁
            emit(cell.paragraphs[0], r[j] if j < len(r) else "", fs, bold=(i == 0))
            if i == 0:
                cell_shade(cell, "D9D9D9")
            tcw = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcw is not None:
                tcw.set(qn("w:type"), "auto"); tcw.set(qn("w:w"), "0")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def image(path, px=None):
    full = os.path.normpath(os.path.join(BASE, path))
    assert os.path.exists(full), full
    width = Cm(17.0 if FAIR else 15.5)
    doc.add_picture(full, width=min(width, Inches(int(px) / 96)) if px else width)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True


def page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, attr in (("w:fldChar", ("w:fldCharType", "begin")), ("w:instrText", None),
                      ("w:fldChar", ("w:fldCharType", "end"))):
        el = OxmlElement(tag)
        if attr:
            el.set(qn(attr[0]), attr[1])
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = "PAGE"
        run._r.append(el)
    style_run(run, 10)


# ---------------- 逐行解析 ----------------
lines = io.open(SRC, encoding="utf-8").read().splitlines()
i = 0
if lines and lines[0].strip() == "---":                      # frontmatter
    i = lines.index("---", 1) + 1
title_block = True
in_refs = False                                              # 「參考文獻」標題之後的段落改為凸排
cover = FAIR and any(l.strip() in ("<!-- 分頁 -->", "<<<分頁>>>") for l in lines)   # 科展封面：分頁標記之前
stats = dict(h=0, p=0, tbl=0, img=0, code=0, quote=0, li=0)
buf = []


def flush():
    global buf
    if buf:
        txt = ""
        for l in buf:
            if txt and not (txt[-1] > "\u2e7f" or l[0] > "\u2e7f"):
                txt += " "
            txt += l
        if txt.startswith("**表") or txt.startswith("**圖"):
            caption(txt)
        else:
            paragraph(txt, center=title_block, size=16 if cover else BODY_PT, hanging=in_refs)
        stats["p"] += 1
        buf = []


while i < len(lines):
    ln = lines[i]
    s = ln.strip()
    if not s or s == "---":
        flush(); i += 1; continue
    if s in ("<!-- 分頁 -->", "<<<分頁>>>"):                   # 封面結束：另起一節，頁碼自 1 起
        flush()
        sec = doc.add_section()
        sec.footer.is_linked_to_previous = False
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:start"), "1")
        sec._sectPr.append(pg)
        cover = False
        i += 1; continue
    if s.startswith("> [!"):                                  # Obsidian callout（導覽），Word 不要
        flush(); i += 1
        while i < len(lines) and lines[i].startswith(">"):
            i += 1
        continue
    m = re.match(r"^(#{1,4})\s+(.*)$", s)
    if m:
        flush()
        lvl, text = len(m.group(1)), m.group(2)
        in_refs = "參考文獻" in text
        if lvl == 1:
            heading(text, 1, center=True)
        elif lvl == 2 and text.startswith("——"):
            heading(text, 3, center=True)                     # 副標
        else:
            if lvl == 2:
                title_block = False
            heading(text, lvl, center=FAIR and lvl == 2)
        stats["h"] += 1; i += 1; continue
    if s.startswith("```"):
        flush(); i += 1; blk = []
        while i < len(lines) and not lines[i].strip().startswith("```"):
            blk.append(lines[i]); i += 1
        code_block(blk); stats["code"] += 1; i += 1; continue
    if s.startswith("|"):
        flush(); rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append(lines[i]); i += 1
        table(rows); stats["tbl"] += 1; continue
    if s.startswith(">"):
        flush(); q = []
        while i < len(lines) and lines[i].strip().startswith(">"):
            q.append(lines[i].strip()[1:].strip()); i += 1
        blockquote(q); stats["quote"] += 1; continue
    m = re.match(r"^!\[\[(.+?)(?:\|(\d+))?\]\]$|^!\[[^\]]*\]\((.+?)\)$", s)
    if m:
        flush(); image(m.group(1) or m.group(3), m.group(2)); stats["img"] += 1; i += 1; continue
    m = re.match(r"^(\d+)\.\s+(.*)$", s)
    if m:
        flush(); list_item(m.group(2), m.group(1) + "."); stats["li"] += 1; i += 1; continue
    if s.startswith("- "):
        flush(); list_item(s[2:], "•"); stats["li"] += 1; i += 1; continue
    buf.append(s); i += 1
flush()
page_number(sec)
doc.save(OUT)
print(f"寫出 {OUT}\n元素統計 {stats}")
