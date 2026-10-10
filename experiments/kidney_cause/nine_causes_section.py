# -*- coding: utf-8 -*-
"""九種病因檢驗之併入段落（計畫書 V3 與科展作品說明書共用）——數字全由 results/nine_causes_nhanes.json 產生，不手打。
    python nine_causes_section.py     → 說明書插入段落（Obsidian 主稿資料夾）＋研究數據表（xlsx）
計畫書由 build_plan_v3.py 直接 import section()。完整文獻與機轉見「腎損傷九種病因與機轉文獻回顧」。"""
import json
import os
import sys

ROOT = os.path.dirname(os.path.abspath(__file__))
R = json.load(open(os.path.join(ROOT, "results", "nine_causes_nhanes.json"), encoding="utf-8"))
PLAN = json.load(open(os.path.join(ROOT, "params", "nine_causes_plan.json"), encoding="utf-8"))
P, PH, SE = R["primary"], R["phenotype"], R["sensitivity"]
FIG = "圖一_九種病因調整盛行率比.png"
DISP = {"B型肝炎": "B 型肝炎", "C型肝炎": "C 型肝炎"}
DIR = {"糖尿病": "代謝", "肥胖": "代謝", "高尿酸血症": "代謝", "痛風": "代謝", "B型肝炎": "感染", "C型肝炎": "感染", "HIV": "感染"}
TYPES = ("只有白蛋白尿", "只有eGFR<60", "兩者皆有")
CNUM = "〇一二三四五六七八九十"
assert all("status" not in P[d] for d in P), "尚有疾病未計算，段落不得產生"
EXPECT = {d: ("腎損傷者較常見" if PLAN["expectations_from_literature"][d].startswith("aPR>1") else "無法區分") for d in P}
assert all(P[d]["verdict"] == EXPECT[d] for d in P), "結果與事前文獻預期不一致，下列敘述須改寫"
HIGHER = sorted((d for d in P if P[d]["verdict"] == "腎損傷者較常見"), key=lambda d: -P[d]["contrasts"]["腎損傷"]["aPR"]["est"])
assert [d for d in P if d not in HIGHER] == ["B型肝炎"], "「只有 B 型肝炎無法區分」之敘述不成立"
for _d in ("高尿酸血症", "痛風"):     # 「集中在 eGFR 下降者」
    _c = PH[_d]["contrasts"]
    assert _c["只有eGFR<60"]["aPR"]["est"] > _c["只有白蛋白尿"]["aPR"]["est"], f"{_d} 集中於 eGFR 下降之敘述不成立"
for _d in ("糖尿病", "C型肝炎", "HIV"):  # 「在白蛋白尿者即已較常見」
    assert PH[_d]["contrasts"]["只有白蛋白尿"]["aPR"]["ci95"][0] > 1, f"{_d} 白蛋白尿之敘述不成立"

pct = lambda p: f"{100 * p:.1f}%" if p >= 0.02 else f"{100 * p:.2f}%"
cnt = lambda x: f"{x:,}"
rr = lambda c: f"{c['est']:.2f}（{c['ci95'][0]:.2f}–{c['ci95'][1]:.2f}）"
apr = lambda d, k="腎損傷", W=P: rr(W[d]["contrasts"][k]["aPR"])
rr2 = lambda c: f"{c['est']:.2f}，95% CI {c['ci95'][0]:.2f}–{c['ci95'][1]:.2f}"      # 句中用，避免括號套括號
apr2 = lambda d, k="腎損傷", W=P: rr2(W[d]["contrasts"][k]["aPR"])
g = lambda d, k: P[d]["groups"][k]
nm = lambda d: DISP.get(d, d)


def cn(i):
    return CNUM[i] if i <= 10 else "十" + CNUM[i - 10] if i < 20 else CNUM[i // 10] + "十" + (CNUM[i % 10] if i % 10 else "")


def label(kind, i, style):
    return f"**{kind}{i}　" if style == "plan" else f"**{kind}{cn(i)}、"


def table_main():
    rows = ["| 方向 | 疾病 | 腎損傷者：陽性／人數；加權盛行率 | 無腎損傷者：陽性／人數；加權盛行率 | 調整盛行率比（95% CI） | 判讀（文獻預期） |",
            "|---|---|---|---|---|---|"]
    for d in P:
        c = lambda k: f"{cnt(g(d, k)['n_events'])}／{cnt(g(d, k)['n'])}；{pct(g(d, k)['prevalence'])}"
        rows.append(f"| {DIR[d]} | {nm(d)} | {c('腎損傷')} | {c('無腎損傷')} | {apr(d)} | {P[d]['verdict']}（符合） |")
    return "\n".join(rows)


def table_pheno():
    rows = ["| 疾病 | 只有白蛋白尿 | 只有 eGFR<60 | 兩者皆有 |", "|---|---|---|---|"]
    rows += [f"| {nm(d)} | " + " | ".join(apr(d, k, PH) for k in TYPES) + " |" for d in P]
    return "\n".join(rows)


def section(fmt, style, n_fig, n_tab, h_meth, h_res):
    """fmt：vault／repo（圖片連結寫法）；style：plan（圖9　）／fair（圖十、）；n_fig、n_tab：本段第一個圖、表號；
    h_meth、h_res：方法與結果小節標題（含編號）。回傳 (方法段, 結果段, 結論一句)。"""
    img = (f"![[figure/nine_causes/{FIG}|620]]" if fmt == "vault" else f"![](../kidney_cause/figures/nine_causes/{FIG})")
    src_t, src_f = ("", "") if style == "plan" else ("（本研究整理）", "（本研究繪製）")
    hv = SE["HIV_20至49歲"]["contrasts"]["腎損傷"]["aPR"]
    meth = f"""{h_meth}

前述三個方向各以一個標籤代表；本段改從臨床角度問：已知腎損傷，可能是哪一種病？依文獻在代謝、免疫、感染三個方向各選三種會傷腎的病——糖尿病、肥胖、高尿酸血症／痛風；紅斑性狼瘡、IgA 腎病變、ANCA 相關血管炎；B 型肝炎、C 型肝炎、HIV，並在計算前提交分析計畫（params/nine_causes_plan.json），寫下每種病依文獻預期的方向。NHANES 1999–2018 量得到其中七項（痛風問卷只有 2007–2018；HIV 只驗 20–49 歲，2009 年起 20–59 歲）；免疫方向三種病沒有診斷問卷、專項抗體或切片，只做文獻回顧。比較腎損傷與無腎損傷成人之加權盛行率，以加權邏輯斯迴歸調整年齡、性別、種族後，以邊際標準化求調整盛行率比，變異以刪一 PSU 摺刀法估計；95% CI 下界大於 1 判讀為「腎損傷者較常見」，上界小於 1 為「較少見」，其餘為「無法區分」。另依腎臟型態（只有白蛋白尿、只有 eGFR<60、兩者皆有）分別比較。
"""
    res = f"""{h_res}

{label('表', n_tab, style)}腎損傷與無腎損傷成人之七項疾病盛行率（NHANES 1999–2018）{src_t}**

{table_main()}

註：陽性／人數為未加權人數，盛行率經抽樣權重加權；調整變項為年齡（十歲一組）、性別與種族；「文獻預期」為計算前寫入分析計畫之方向。

{img}

{label('圖', n_fig, style)}腎損傷成人之調整盛行率比（左：腎損傷對無腎損傷；右：依腎臟型態對兩項皆正常者；橫線為 95% CI）{src_f}**

調整年齡、性別與種族後，腎損傷成人較常見的依序為：{'、'.join(f'{nm(d)}（{P[d]["contrasts"]["腎損傷"]["aPR"]["est"]:.2f}）' for d in HIGHER)}；B 型肝炎無法區分（{apr2('B型肝炎')}）。{len(P)} 項皆與計算前寫下的文獻預期方向一致。以盛行率而言，腎損傷成人中最常見的是肥胖（{pct(g('肥胖', '腎損傷')['prevalence'])}）、高尿酸血症（{pct(g('高尿酸血症', '腎損傷')['prevalence'])}）與糖尿病（{pct(g('糖尿病', '腎損傷')['prevalence'])}）；感染方向三種病都不到 2%。HIV 陽性人數少（腎損傷者 {g('HIV', '腎損傷')['n_events']} 人），限 20–49 歲之敏感度分析結果相近（{rr2(hv)}）。

{label('表', n_tab + 1, style)}依腎臟型態之調整盛行率比（對照組：eGFR 與 ACR 皆正常）{src_t}**

{table_pheno()}

腎損傷的型態本身就是線索：高尿酸血症與痛風集中在 eGFR 下降者（高尿酸血症：只有 eGFR<60 為 {apr2('高尿酸血症', '只有eGFR<60', PH)}；只有白蛋白尿為 {PH['高尿酸血症']['contrasts']['只有白蛋白尿']['aPR']['est']:.2f}），與「過濾減少使尿酸排泄減少」一致，提醒尿酸較可能是腎損傷的結果；糖尿病、C 型肝炎與 HIV 在白蛋白尿者即已較常見（只有白蛋白尿：糖尿病 {apr2('糖尿病', '只有白蛋白尿', PH)}；C 型肝炎 {apr2('C型肝炎', '只有白蛋白尿', PH)}；HIV {apr2('HIV', '只有白蛋白尿', PH)}），與這些病以腎絲球病變傷腎一致。這是橫斷面共存，不代表因果方向。
"""
    concl = (f"已知腎損傷時，依公開資料可先想代謝方向（糖尿病 {P['糖尿病']['contrasts']['腎損傷']['aPR']['est']:.2f} 倍、痛風、高尿酸與肥胖），"
             f"感染方向中 C 型肝炎與 HIV 亦較常見而 B 型肝炎無法區分；腎損傷的型態（白蛋白尿或 eGFR 下降）可進一步縮小範圍。"
             "免疫方向三種病在公開資料量不到，須靠專項抗體與腎臟切片。")
    return meth, res, concl


def data_rows():
    """研究數據表：每病×每比較一列。"""
    out = []
    for scope, W, keys in (("主要：腎損傷對無腎損傷", P, ("腎損傷",)), ("依腎臟型態", PH, TYPES), ("敏感度：HIV 限 20–49 歲", SE, ("腎損傷",))):
        for d in W:
            for k in keys:
                grp = W[d]["groups"]
                ref = "無腎損傷" if "無腎損傷" in grp else "無異常"
                c = W[d]["contrasts"][k]
                out.append([scope, nm(d.replace("HIV_20至49歲", "HIV")), DIR.get(d, "感染"), k, grp[k]["n_events"], grp[k]["n"],
                            round(grp[k]["prevalence"], 5), grp[ref]["n_events"], grp[ref]["n"], round(grp[ref]["prevalence"], 5),
                            round(c["PR_crude"]["est"], 3), round(c["aPR"]["est"], 3), round(c["aPR"]["ci95"][0], 3),
                            round(c["aPR"]["ci95"][1], 3), round(c["aOR"]["est"], 3), W[d].get("verdict", "")])
    return out


def excel(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "九種病因檢驗"
    head = ["分析", "疾病", "方向", "比較組", "比較組陽性", "比較組人數", "比較組加權盛行率", "對照組陽性", "對照組人數",
            "對照組加權盛行率", "粗盛行率比", "調整盛行率比", "95% CI 下界", "95% CI 上界", "調整勝算比", "判讀"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in data_rows():
        ws.append(r)
    ws2 = wb.create_sheet("說明")
    for line in ("資料：NHANES 1999–2018 成人（≥20 歲），腎臟狀態可判定者；腎損傷＝eGFR<60 或 ACR≥30。",
                 "人數為未加權；盛行率經 w_mec20 加權；調整盛行率比調整年齡（十歲一組）、性別、種族，邊際標準化。",
                 "95% CI 以刪一 PSU 摺刀法；對照組：主要分析＝無腎損傷，依腎臟型態＝eGFR 與 ACR 皆正常。",
                 f"分析計畫 params/nine_causes_plan.json（SHA256 {R['plan_sha256'][:16]}…）；結果 results/nine_causes_nhanes.json。",
                 "紅斑性狼瘡、IgA 腎病變、ANCA 相關血管炎：NHANES 無可用資料，未檢驗。"):
        ws2.append([line])
    for col, w in zip("ABCDEFGHIJKLMNOP", (22, 12, 6, 14, 10, 10, 14, 10, 10, 14, 10, 12, 11, 11, 10, 14)):
        ws.column_dimensions[col].width = w
    ws2.column_dimensions["A"].width = 110
    wb.save(path)
    print(f"[寫出] {path}（{ws.max_row - 1} 列）")


def main():
    sys.path.insert(0, ROOT)
    import build_fair_v3_2 as F
    out = sys.argv[1] if len(sys.argv) > 1 else F.VAULT
    led = json.load(open(os.path.join(ROOT, "results", "provenance.json"), encoding="utf-8"))["files"]
    n_all, mb = len(led), sum(v["bytes"] for v in led.values()) / 1e6
    n_dev = n_all - sum(1 for k in led if k.endswith("_L.xpt"))      # 2021–2023 檔名以 _L 結尾
    meth, res, concl = section("vault", "fair", 10, 19, "#### （十一）九種病因之公開資料檢驗", "### 八、已知腎損傷，可能是哪一種病？")
    txt = ("# 科展作品說明書插入段落：九種病因之公開資料檢驗\n\n"
           "> [!info] 由 `experiments/kidney_cause/nine_causes_section.py` 自結果檔產生。表號、圖號接在說明書原有之表十八、圖九之後；"
           "若說明書表號不同請自行調整。不修改說明書原有文字。\n\n"
           f"另需改一處數字：「貳、一、資料來源」表中開發資料檔案數改為 {n_dev}、合計改為 {n_all} 個檔案（約 {mb:.0f} MB），因出處帳本新增 HIV 檢驗 10 檔。\n\n"
           "## 插入「參、二、研究方法」之末\n\n" + meth +
           "\n## 插入「肆、研究結果」之末（七、把回推結果做成工具之後）\n\n" + res +
           "\n## 插入「陸、結論」之末\n\n" + concl + "\n")
    p = os.path.join(out, "科展作品說明書_插入段落_九種病因.md")
    open(p, "w", encoding="utf-8").write(txt)
    print(f"[寫出] {p}（{len(txt):,} 字元）")
    excel(os.path.join(out, "研究數據表_九種病因.xlsx"))


if __name__ == "__main__":
    main()
