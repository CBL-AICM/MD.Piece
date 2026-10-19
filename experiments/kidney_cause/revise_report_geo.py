# -*- coding: utf-8 -*-
"""研究報告（整合版 Word）全篇納入 NCBI GEO 腎臟切片轉錄體（免疫方向補充）。
數字一律讀 results/immune_geo_batchsafe.json（修正一）與 results/immune_geo_amendment2.json（修正二），
敘述依賴的結論以 assert 鎖住；結論變了程式會停，不會寫出與數字矛盾的句子。
每處修改以原文開頭定位（找不到或找到多處即停）。原檔不動，另存新檔。
    python -X utf8 revise_report_geo.py <輸入.docx（已換圖）> <輸出.docx>"""
import copy
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm
from PIL import Image

ROOT = Path(__file__).resolve().parent
J = lambda n: json.load(open(ROOT / "results" / n, encoding="utf-8"))
B = J("immune_geo_batchsafe.json")
C = J("immune_geo_amendment2.json") if (ROOT / "results" / "immune_geo_amendment2.json").exists() else None
TODO = "【待補：修正二計算中】"
src, dst = sys.argv[1:3]
assert Path(src).resolve() != Path(dst).resolve(), "不覆寫原檔"
d = Document(src)

f2 = lambda x: f"{x:.2f}"
f3 = lambda x: f"{x:.3f}"
ci = lambda c: f"{c[0]:.2f}–{c[1]:.2f}"
pv = lambda p: "<0.001" if p < 0.001 else f"{p:.3f}"
b1, b2, b3 = B["B1_primary"], B["B2_immune_vs_metabolic"], B["B3_SLE"]
DB = B["decision"]
P_, G_ = "PANEL_LR", "GENE_LR"


# ---------- 段落工具 ----------
def para(prefix):
    hits = [p for p in d.paragraphs if p.text.startswith(prefix)]
    assert len(hits) == 1, f"定位失敗（{len(hits)} 處）：{prefix}"
    return hits[0]


def set_text(p, text):
    """保留第一個 run 的字型，其餘 run 刪除。"""
    runs = p.runs
    assert runs, p.text
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    return p


def replace_in(prefix, old, new):
    p = para(prefix)
    assert p.text.count(old) == 1, f"原句不唯一或不存在：{old}"
    return set_text(p, p.text.replace(old, new))


def insert_after(anchor, like, text):
    """在 anchor 之後插入一段，格式複製自 like。"""
    el = copy.deepcopy(like._p)
    anchor._p.addnext(el)
    from docx.text.paragraph import Paragraph
    p = Paragraph(el, anchor._parent)
    return set_text(p, text) if text is not None else p


def cell(table, r, c, text):
    p = table.rows[r].cells[c].paragraphs[0]
    set_text(p, text)
    for extra in table.rows[r].cells[c].paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def add_row(table, values, like_row=-1):
    tr = copy.deepcopy(table.rows[like_row]._tr)
    table._tbl.append(tr)
    for i, v in enumerate(values):
        cell(table, len(table.rows) - 1, i, v)


def table_after(prefix):
    p = para(prefix)
    for el in p._p.itersiblings():
        if el.tag.endswith("}tbl"):
            return next(t for t in d.tables if t._tbl is el)
    raise AssertionError(prefix)


BODY = para("合計 285 個檔案")          # 一般內文段落格式
H_METHOD = para("十一、九種病因之公開資料檢驗")
H_RESULT = para("十一、腎損傷成人中各候選病因的盛行率")
LIST = para("8. 建立網頁工具")

PRE = J("immune_geo.json")
REP = B["B4_replication"]
pp1, pp3 = b1["per_panel"], b3["per_panel"]
ifn3 = pp3["interferon"]["auroc"]
rs3 = b3["random_gene_sets"]
# 敘述依賴的結論（修正一）
assert all(DB[k]["回推成立"] for k in ["B1_primary", "B2_immune_vs_metabolic", "B3_SLE"])
assert not DB["B2_immune_vs_metabolic"]["PANEL_LR"] and DB["B2_immune_vs_metabolic"]["GENE_LR"]
assert DB["免疫路徑專屬"]["B3_SLE"] and max(pp3, key=lambda k: pp3[k]["auroc"]) == "interferon"
assert sorted(pp1, key=lambda k: -pp1[k]["auroc"])[:3] == ["myeloid", "complement", "neutrophil"]
assert all(REP[k][m]["auroc_mean"] > .5 for k in REP for m in [P_, G_])
assert not PRE["decision"]["回推成立"]

import immune_geo as ig   # noqa: E402
geo_n = {s: sum(ig.load(s, p)[1].shape[0] for p in ["GPL22945", "GPL24120"]) for s in ["GSE104948", "GSE104954"]}
n_gene = ig.load("GSE104948", "GPL22945")[1].shape[1]
assert n_gene == ig.load("GSE104948", "GPL24120")[1].shape[1]

# ---------- 摘要 ----------
replace_in("腎損傷的病因大致可分為", "免疫方向（抗核抗體）僅0.546，低於只用年齡與性別的0.596。",
           "免疫方向（抗核抗體）僅0.546，低於只用年齡與性別的0.596；改以美國國家衛生研究院公開的病理確診腎臟切片轉錄體分析，"
           f"在同一處理批次內，狼瘡腎炎與ANCA血管炎可與非免疫腎病區分（AUROC {f2(b3[P_]['auroc_mean'])}與{f2(b1[G_]['auroc_mean'])}）。")
replace_in("腎損傷的病因大致可分為", "常規檢驗可回推代謝與部分感染方向，但無法回推免疫方向，也不能取代臨床診斷。",
           "常規檢驗可回推代謝與部分感染方向；免疫方向的訊號在腎臟組織中，常規檢驗測不到，仍須靠專項抗體與切片。本研究不能取代臨床診斷。")
replace_in("Kidney damage may stem", "below the 0.596 from age and sex alone.",
           "below the 0.596 from age and sex alone. In public kidney biopsy transcriptomes from the NIH Gene Expression Omnibus, "
           "compared within the same processing batch, lupus nephritis and ANCA vasculitis were separable from non-immune kidney "
           f"disease (AUROC {f2(b3[P_]['auroc_mean'])} and {f2(b1[G_]['auroc_mean'])}).")
replace_in("Kidney damage may stem", "Routine tests can trace metabolic and some infectious causes but not immune causes, and they cannot replace clinical diagnosis.",
           "Routine tests can trace metabolic and some infectious causes. The immune signal lies in kidney tissue, which routine "
           "tests do not measure, so immune causes still need specific antibodies and biopsy. None of this replaces clinical diagnosis.")

# ---------- 壹、貳 ----------
replace_in("依 KDIGO 2024 指引", "（Yang et al., 2024）。本研究的標籤",
           "（Yang et al., 2024）。歐洲腎臟 cDNA 庫另公開了 ANCA 相關血管炎、狼瘡腎炎等病理確診腎病的切片轉錄體，"
           "可比較不同腎病在組織中的基因表現（Grayson et al., 2018）。本研究的標籤")
replace_in("本研究旨在探討腎損傷時", "研究中只使用公開、去識別化的 NHANES 資料。",
           "研究中只使用公開、去識別化的資料。主要資料為 NHANES；NHANES 沒有免疫性腎病的診斷，免疫方向另以美國國家衛生研究院"
           "國家生物技術資訊中心（NCBI）Gene Expression Omnibus（GEO）公開的腎臟切片轉錄體補充。")
cell(table_after("本研究旨在探討腎損傷時"), 3, 2, "1999–2004 年剩餘血清次樣本；另以 GEO 病理確診切片轉錄體補充")
replace_in("除以三個標籤評估模型之外", "以公開資料比較其在腎損傷與無腎損傷成人中的盛行率，",
           "以公開資料比較其在腎損傷與無腎損傷成人中的盛行率（免疫方向三種病在 NHANES 沒有資料，改以切片轉錄體比較組織訊號），")
p = para("此外，本研究以探索性暴露分析")
set_text(p, p.text + "免疫方向在 NHANES 只有抗核抗體可用，因此另以 GEO 公開的病理確診切片轉錄體，在同一處理批次內比較免疫、代謝與兩者皆非的腎病，"
         "檢視免疫方向的訊號是否存在於腎臟組織。")
p = para("本研究事前設定兩項假說")
set_text(p, p.text + "切片轉錄體分析是在 NHANES 結果完成後才加入的補充，不屬於上述兩項事前假說；其分析計畫與兩次修正的判定規則，皆在計算前提交（肆、十二）。")

# ---------- 參、資料來源與品質管控 ----------
p = para("合計 285 個檔案")
set_text(p, "NHANES " + p.text)
add_row(table_after("美國國家健康與營養調查（NHANES）由"), ["切片轉錄體（GEO）", "2018 年公開", "5", "免疫方向補充"])
insert_after(p, p, f"GEO 由美國國家衛生研究院國家醫學圖書館之 NCBI 維護，存放研究者公開的基因表現資料。本研究使用歐洲腎臟 cDNA 庫的腎臟切片轉錄體，"
             f"GSE104948 為腎絲球（{geo_n['GSE104948']} 例）、GSE104954 為腎小管間質（{geo_n['GSE104954']} 例），各含兩種晶片平台"
             "（Eddy et al., 2018a, 2018b），另下載 NCBI 基因資訊檔對照基因名稱。每例診斷皆經腎臟病理確認，對照組為活體捐腎者與腫瘤腎切除檢體。"
             "5 個檔案同樣記錄網址與 SHA256；資料已去識別，本研究不接觸任何受試者。")
t = table_after("三、研究品質管控")
add_row(t, ["處理批次核對", "依樣本來源的批次代碼建表，只比較同一批次內的疾病", "把批次差異誤當成疾病差異"])
add_row(t, ["獨立重算", "以不引用主程式的程式，從原始檔重算切片轉錄體修正一的各項比較", "程式算錯"])

# ---------- 肆、方法 ----------
insert_after(LIST, LIST, "9. 以 GEO 腎臟切片轉錄體補充免疫方向，只比較同一處理批次內的疾病。")
replace_in("前述三個方向各以一個標籤代表", "免疫方向三種病沒有診斷問卷、專項抗體或切片，只做文獻回顧。",
           "免疫方向三種病在 NHANES 沒有診斷問卷、專項抗體或切片，以文獻回顧與切片轉錄體（肆、十二）補充。")
M = [
 "NHANES 沒有免疫性腎病的診斷，抗核抗體陽性也不代表腎臟受到免疫攻擊。因此本研究另以 GEO 的病理確診切片轉錄體檢視免疫方向（參、一）。"
 "分析計畫在計算前提交至版本控制，任務、模型、評估方式與判定規則均事前寫定。",
 "下載後核對樣本來源名稱，發現名稱中的 H 代碼為處理批次，且疾病與批次幾乎完全重疊。糖尿病腎病變在 H1 與 H7，高血壓腎病變只在 H2，"
 "IgA 腎病變只在 H4，狼瘡腎炎、膜性腎病變與局部節段性腎絲球硬化（FSGS）在 H5，ANCA 相關血管炎在 H7。GEO 的資料處理說明寫明，"
 "各批次分別正規化後合併，再以 ComBat 做批次校正（Johnson et al., 2007）；批次與疾病重疊時，這一步會連同大部分疾病差異一起移除。"
 "原計畫的跨批次比較因此失效，結果照原樣保存而不解讀。修正一改為只比較同一批次內的疾病；獨立稽核之後，另將代謝方向本身與狼瘡腎炎對膜性腎病變"
 "兩項比較列為修正二。兩次修正的判定規則都在計算前提交，但都訂於得知原計畫結果之後，屬事後分析。",
 "同一批次內共有六項比較。免疫對非免疫為 ANCA 相關血管炎對糖尿病腎病變、FSGS 與微小變化病（H7），以及狼瘡腎炎對 FSGS（H5）；"
 "免疫對代謝為 ANCA 相關血管炎對糖尿病腎病變（H7）；代謝對兩者皆非為糖尿病腎病變對 FSGS 與微小變化病（H7），以及糖尿病腎病變對"
 "微小變化病與薄基底膜病（H1）；免疫對免疫為狼瘡腎炎對膜性腎病變（H5）。IgA 腎病變與高血壓腎病變各自只在一個批次，同批次內沒有對照，無法評估。",
 "模型有兩種。免疫路徑模型以事前選定的六條路徑（干擾素、補體、B 細胞與漿細胞、T 細胞、單核巨噬細胞、嗜中性球）的分數作邏輯迴歸；"
 f"全基因模型以 {n_gene:,} 個基因作強正則化邏輯迴歸。超參數固定，不調參。評估採分層五折、重複 20 次，標準化只用訓練折，"
 "95% CI 以第一次重複 bootstrap 2,000 次估計，另把標籤打亂 1,000 次作置換檢定。每項比較須同時符合三個條件才判為可區分，"
 "即交叉驗證 AUROC 的 95% CI 下界大於 0.5、置換 p 小於 0.025（兩個模型以 Bonferroni 校正），以及在腎小管間質重做時同一模型 AUROC 大於 0.5。"
 "狼瘡腎炎對膜性腎病變另檢定干擾素路徑單獨的 AUROC，95% CI 下界大於 0.5、置換 p 小於 0.05 且腎小管間質大於 0.5 時，判定干擾素訊號為狼瘡特有。"
 "腎小管間質與腎絲球取自同一批切片，只代表另一個腎臟區域的重現，不是獨立樣本。",
]
body_m = para("文獻部分以 PubMed 檢索")
h = insert_after(body_m, H_METHOD, "十二、腎臟切片轉錄體（免疫方向補充）")
for t_ in M:
    h = insert_after(h, body_m, t_)

# ---------- 伍、結果 ----------
set_text(para("由於 NHANES 沒有免疫性腎病的診斷資料"), "NHANES 沒有免疫性腎病的診斷資料，以病理確診切片轉錄體所做的補充分析見伍、十三。")
rep = lambda k: f"{f2(REP[k][P_]['auroc_mean'])}／{f2(REP[k][G_]['auroc_mean'])}"
zh = dict(interferon="干擾素", complement="補體", b_plasma="B 細胞與漿細胞", t_cell="T 細胞", myeloid="單核巨噬細胞", neutrophil="嗜中性球")
R = [
 f"在同一處理批次內，免疫性腎病與非免疫腎病可以區分（圖 10、表 10）。ANCA 相關血管炎對糖尿病腎病變、FSGS 與微小變化病，"
 f"全基因模型 AUROC 為 {f3(b1[G_]['auroc_mean'])}（95% CI {ci(b1[G_]['auroc_ci_rep0'])}，置換 p {pv(b1[G_]['permutation_p'])}），"
 f"免疫路徑模型為 {f3(b1[P_]['auroc_mean'])}（95% CI {ci(b1[P_]['auroc_ci_rep0'])}）；狼瘡腎炎對 FSGS，免疫路徑模型為 "
 f"{f3(b3[P_]['auroc_mean'])}（95% CI {ci(b3[P_]['auroc_ci_rep0'])}，置換 p {pv(b3[P_]['permutation_p'])}），全基因模型為 "
 f"{f3(b3[G_]['auroc_mean'])}。兩項在腎小管間質皆重現（路徑／全基因 AUROC 分別為 {rep('B1')} 與 {rep('B3')}）。",
 f"免疫對代謝只有 ANCA 相關血管炎對糖尿病腎病變一項（{b2['n_pos']} 對 {b2['n'] - b2['n_pos']} 例）。全基因模型 AUROC 為 "
 f"{f3(b2[G_]['auroc_mean'])}（95% CI {ci(b2[G_]['auroc_ci_rep0'])}），免疫路徑模型只有 {f3(b2[P_]['auroc_mean'])}，未達顯著；"
 "也就是說，事前選的六條免疫路徑分不開這兩種病，全基因模型的判別力主要來自路徑以外的基因。原計畫對同一比較只檢定免疫路徑模型而判為不成立，"
 "修正一「任一模型成立即可」的規則是在已知全基因結果後才寫入，所以這一項屬事後分析，不能當成事前確認。",
 None,   # 代謝方向本身（修正二）
 None,   # 狼瘡腎炎對膜性腎病變（修正二）
 f"各路徑單獨來看，ANCA 相關血管炎以單核巨噬細胞（{f2(pp1['myeloid']['auroc'])}）、補體（{f2(pp1['complement']['auroc'])}）"
 f"與嗜中性球（{f2(pp1['neutrophil']['auroc'])}）較高；狼瘡腎炎對 FSGS 以干擾素最高（{f2(ifn3)}），六條路徑合起來的模型"
 f"（第一次重複 {f2(rs3['panel_auroc_rep0'])}）高於 1,000 組同大小隨機基因集的第 95 百分位（{f2(rs3['p95'])}）。"
 "各路徑單獨的 AUROC 未經交叉驗證與多重比較校正，只作描述。",
 f"原計畫的跨批次比較（狼瘡腎炎與 IgA 腎病變對糖尿病與高血壓腎病變）因批次與疾病重疊而失效。其全基因模型交叉驗證 AUROC 只有 "
 f"{f3(PRE['T1_primary']['GENE_LR']['auroc_mean'])}，保留測試集 {f3(PRE['T1_primary']['holdout']['GENE_LR']['auroc'])}，"
 "訓練集卻幾乎全對，屬於記憶批次的表現；結果檔照原樣保存，數字不引用。",
]

# ---------- 修正二結果（代謝方向本身、狼瘡腎炎對膜性腎病變）----------
assert C is not None, "修正二結果檔不存在"
c1, c2, c3 = C["C1_metabolic_H7"], C["C2_metabolic_H1"], C["C3_SLE_vs_MN"]
DC, it = C["decision"], C["interferon_test"]
R2 = C["C4_replication"]
# 敘述依賴的結論（修正二）：結論變了就停
assert not DC["C1_metabolic_H7"]["回推成立"] and not DC["C2_metabolic_H1"]["回推成立"]
assert all(c[m]["permutation_p"] > .025 for c in (c1, c2) for m in (P_, G_))
assert DC["C3_SLE_vs_MN"]["回推成立"] and DC["干擾素專屬狼瘡_C3"] and DC["免疫路徑專屬_C3"] and DC["描述_干擾素為最高路徑_C3"]
assert not any(DC["批次記憶警示"].values())
for c in (c1, c2):   # 兩批糖尿病腎病變：補體、單核巨噬、T 細胞較高，干擾素較低
    assert all(c["per_panel"][k]["mean_diff_z"] > 0 for k in ("complement", "myeloid", "t_cell"))
    assert c["per_panel"]["interferon"]["mean_diff_z"] < 0
lo_dn = min(c1["n_pos"], c2["n_pos"]); hi_dn = max(c1["n_pos"], c2["n_pos"])
pts = [c1[P_]["auroc_mean"], c1[G_]["auroc_mean"], c2[P_]["auroc_mean"], c2[G_]["auroc_mean"]]
reps_dn = [R2[k][m]["auroc_mean"] for k in ("C1", "C2") for m in (P_, G_)]
pp_c1 = c1["per_panel"]
R[2] = (f"代謝方向本身沒有達到判定。糖尿病腎病變對 FSGS 與微小變化病（H7，{c1['n_pos']} 對 {c1['n'] - c1['n_pos']} 例），"
        f"免疫路徑模型 AUROC {f3(c1[P_]['auroc_mean'])}（95% CI {ci(c1[P_]['auroc_ci_rep0'])}，置換 p {pv(c1[P_]['permutation_p'])}），"
        f"全基因模型 {f3(c1[G_]['auroc_mean'])}（95% CI {ci(c1[G_]['auroc_ci_rep0'])}，置換 p {pv(c1[G_]['permutation_p'])}）；"
        f"糖尿病腎病變對微小變化病與薄基底膜病（H1，{c2['n_pos']} 對 {c2['n'] - c2['n_pos']} 例），分別為 {f3(c2[P_]['auroc_mean'])}"
        f"（95% CI {ci(c2[P_]['auroc_ci_rep0'])}，置換 p {pv(c2[P_]['permutation_p'])}）與 {f3(c2[G_]['auroc_mean'])}"
        f"（95% CI {ci(c2[G_]['auroc_ci_rep0'])}）。點估計介於 {f2(min(pts))} 至 {f2(max(pts))}，腎小管間質重做為 "
        f"{f2(min(reps_dn))} 至 {f2(max(reps_dn))}，但糖尿病腎病變每批只有 {lo_dn} 至 {hi_dn} 例，置換 p 都大於 0.025，"
        "依規則不能判為可區分；這是例數不足下的「無法確認」，不是「確認沒有差異」。兩批的糖尿病腎病變都以補體、單核巨噬細胞與 T 細胞路徑較高"
        f"（H7 單路徑 AUROC {f2(pp_c1['complement']['auroc'])}、{f2(pp_c1['myeloid']['auroc'])}、{f2(pp_c1['t_cell']['auroc'])}），"
        f"干擾素路徑反而較低（{f2(pp_c1['interferon']['auroc'])}）。")
R[3] = (f"狼瘡腎炎對膜性腎病變（{c3['n_pos']} 對 {c3['n'] - c3['n_pos']} 例）可以區分。免疫路徑模型 AUROC {f3(c3[P_]['auroc_mean'])}"
        f"（95% CI {ci(c3[P_]['auroc_ci_rep0'])}，置換 p {pv(c3[P_]['permutation_p'])}），全基因模型 {f3(c3[G_]['auroc_mean'])}；"
        f"六條路徑合起來的模型（第一次重複 {f2(c3['random_gene_sets']['panel_auroc_rep0'])}）高於 1,000 組隨機基因集的第 95 百分位"
        f"（{f2(c3['random_gene_sets']['p95'])}）。干擾素路徑單獨即達 {f3(it['glom']['auroc'])}（95% CI {ci(it['glom']['ci'])}，"
        f"置換 p {pv(it['glom']['permutation_p'])}），腎小管間質為 {f3(it['tubint']['auroc'])}，符合計算前寫定的「干擾素訊號為狼瘡特有」判定。"
        "兩者同屬免疫複合體沉積的腎病，狼瘡腎炎多出的主要是干擾素訊號。")
C_DISC = (f"代謝方向本身（糖尿病腎病變對既非免疫也非代謝的腎病）點估計 {f2(min(pts))} 至 {f2(max(pts))}，但例數太少而未達判定；"
          f"免疫病之間，狼瘡腎炎可由干擾素路徑與膜性腎病變分開（{f2(it['glom']['auroc'])}）")
C_CONC = (f"，狼瘡腎炎與同為免疫複合體腎病的膜性腎病變也能以干擾素路徑區分（{f2(it['glom']['auroc'])}）；"
          f"糖尿病腎病變對非免疫腎病的點估計為 {f2(min(pts))} 至 {f2(max(pts))}，但每批只有 {lo_dn} 至 {hi_dn} 例，未達判定")

# ---------- 表 9（九種病因）免疫三列、表註、文獻段 ----------
t9 = table_after("表 9　九種候選病因")
rows9 = {r.cells[1].text: i for i, r in enumerate(t9.rows)}
assert all(t9.rows[rows9[k]].cells[4].text == "NHANES 無資料" for k in ["紅斑性狼瘡", "IgA 腎病變", "ANCA 相關血管炎"])
cell(t9, rows9["紅斑性狼瘡"], 4, f"NHANES 無資料；切片轉錄體對 FSGS 之 AUROC {f2(b3[P_]['auroc_mean'])}（免疫路徑模型），干擾素路徑最高（{f2(ifn3)}）")
cell(t9, rows9["IgA 腎病變"], 4, "NHANES 無資料；切片轉錄體只有單一處理批次，無法評估")
cell(t9, rows9["ANCA 相關血管炎"], 4, f"NHANES 無資料；切片轉錄體對非免疫腎病之 AUROC {f2(b1[G_]['auroc_mean'])}（全基因模型）")
p = para("註　調整盛行率比為腎損傷對無腎損傷")
set_text(p, p.text + "切片轉錄體之數字為同一處理批次內、腎絲球之交叉驗證 AUROC（伍、十三）。")
p = para("免疫方向的三種病在 NHANES 沒有可用資料")
set_text(p, p.text + "切片轉錄體的結果與上述機轉相符，狼瘡腎炎以干擾素路徑升高為主，ANCA 相關血管炎以單核巨噬細胞、補體與嗜中性球路徑升高為主（伍、十三）。")

# ---------- 陸、討論 ----------
p = para("免疫方向的假說一不成立")
set_text(p, p.text + "切片轉錄體顯示免疫性腎病在腎臟組織中可與非免疫腎病區分（伍、十三），所以常規檢驗回推不了免疫方向，較可能是常規項目中"
         "沒有反映這些路徑的檢驗，而不是免疫性腎病缺乏可辨識的特徵。")
replace_in("現有公開資料可支持以下結論", "免疫方向則未見足夠訊號。", "免疫方向在常規檢驗中未見足夠訊號，但在腎臟組織中可以區分。")
replace_in("8. 九種病因之比較為單次橫斷面資料", "免疫方向三種病無法量測", "免疫方向三種病在 NHANES 無法量測")
lim8 = para("8. 九種病因之比較為單次橫斷面資料")
insert_after(lim8, lim8, "9. 切片轉錄體是確診後的組織，不是常規抽血驗尿。疾病與處理批次幾乎完全重疊，只能在同一批次內比較；兩次修正皆屬事後分析，"
             "跨比較也未校正多重檢定。各組例數少（糖尿病腎病變每批 5 至 7 例），信賴區間寬；ERCB 的糖尿病腎病變是臨床上需要切片的非典型個案，"
             "不代表一般糖尿病腎病變。IgA 腎病變與高血壓腎病變無法評估，也沒有獨立的外部資料集。")
p = para("近期將以未參與模型開發或重新校準的獨立資料")
set_text(p, p.text + "免疫方向則可評估切片中區分疾病的干擾素與補體路徑，能否以血液或尿液檢驗測得。")

# ---------- 參考文獻（PubMed 核對：PMID 16632515、28262777、28760771）----------
def add_ref(after_prefix, plain, italic, tail):
    ref = para(after_prefix)
    q = insert_after(ref, ref, plain)
    r0 = q.runs[0]
    for txt, itl in [(italic, True), (tail, False)]:
        r = q.add_run(txt)
        if r0._r.rPr is not None:
            r._r.insert(0, copy.deepcopy(r0._r.rPr))
        r.italic = itl


add_ref("Johnson, R. J., Gretch", "Johnson, W. E., Li, C., & Rabinovic, A. (2007). Adjusting batch effects in microarray expression data "
        "using empirical Bayes methods. ", "Biostatistics, 8", "(1), 118–127. https://doi.org/10.1093/biostatistics/kxj037")
add_ref("Fabrizi, F.", "Flyvbjerg, A. (2017). The role of the complement system in diabetic nephropathy. ", "Nature Reviews Nephrology, 13",
        "(5), 311–318. https://doi.org/10.1038/nrneph.2017.31")
add_ref("Tervaert, T. W.", "Tesch, G. H. (2017). Diabetic nephropathy – is this an immune disorder? ", "Clinical Science, 131",
        "(16), 2183–2199. https://doi.org/10.1042/CS20160636")
assert "Johnson et al., 2007" in para("下載後核對樣本來源名稱").text

# ---------- 伍、十三（新節）＋圖 10＋表 10 ----------
END12 = para("免疫方向的三種病在 NHANES 沒有可用資料")
body_r = para("調整年齡、性別與種族後，腎損傷成人較常見的依序為")
x = insert_after(END12, H_RESULT, "十三、腎臟切片轉錄體中的免疫、代謝與兩者皆非")
for t_ in R:
    x = insert_after(x, body_r, t_ if t_ is not None else TODO)
fig10 = ROOT / "figures" / "report" / "圖10_免疫方向切片轉錄體.png"
img9 = para("圖 9 腎損傷成人之調整盛行率比")
if C is not None and fig10.exists():
    pic = insert_after(x, d.paragraphs[[q._p for q in d.paragraphs].index(img9._p) - 1], None)
    rid, _ = d.part.get_or_add_image(str(fig10))
    pic._p.xpath(".//a:blip")[0].set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rid)
    w, h_ = Image.open(fig10).size
    for e in pic._p.xpath(".//wp:extent") + pic._p.xpath(".//a:ext"):
        e.set("cx", str(Cm(17))); e.set("cy", str(round(Cm(17) * h_ / w)))
    for e in pic._p.xpath(".//wp:docPr"):
        e.set("id", "1010")
    x = insert_after(pic, img9, "圖 10 腎臟切片轉錄體同一處理批次內之比較（(a) 點為 20 次重複平均，橫線為第一次重複之 95% CI，"
                     "實心為腎絲球、空心為腎小管間質；(b) 各免疫路徑單獨之 AUROC，腎絲球）")
else:
    x = insert_after(x, body_r, "【待補：圖 10（修正二完成後產生）】")
cap4 = para("表 4 肝炎軸分型")
x = insert_after(x, cap4, "表 10 腎臟切片轉錄體同一處理批次內之比較（腎絲球，20 次重複平均）")
t10 = copy.deepcopy(table_after("表 4 肝炎軸分型")._tbl)
x._p.addnext(t10)
t10 = next(t for t in d.tables if t._tbl is t10)
while len(t10.rows) > 2:
    t10._tbl.remove(t10.rows[-1]._tr)
cell_ = lambda r: f"{f3(r['auroc_mean'])}（{ci(r['auroc_ci_rep0'])}）"
for i, v in enumerate(["方向", "比較（批次）", "陽性對對照", "免疫路徑模型", "全基因模型", "腎小管間質（路徑／全基因）", "可區分"]):
    cell(t10, 0, i, v)
ROWS10 = [("免疫對非免疫", "ANCA 對 DN、FSGS、MCD（H7）", b1, "B1", DB["B1_primary"]),
          ("免疫對非免疫", "SLE 對 FSGS（H5）", b3, "B3", DB["B3_SLE"]),
          ("免疫對代謝", "ANCA 對 DN（H7）", b2, "B2", DB["B2_immune_vs_metabolic"])]
if C is not None:
    ROWS10 += [("代謝對兩者皆非", "DN 對 FSGS、MCD（H7）", C["C1_metabolic_H7"], "C1", C["decision"]["C1_metabolic_H7"]),
               ("代謝對兩者皆非", "DN 對 MCD、TMD（H1）", C["C2_metabolic_H1"], "C2", C["decision"]["C2_metabolic_H1"]),
               ("免疫對免疫", "SLE 對 MN（H5）", C["C3_SLE_vs_MN"], "C3", C["decision"]["C3_SLE_vs_MN"])]
    REP.update(C["C4_replication"])
for k, (dr, lab, r, code, dec) in enumerate(ROWS10):
    vals = [dr, lab, f"{r['n_pos']} 對 {r['n'] - r['n_pos']}", cell_(r[P_]), cell_(r[G_]), rep(code), ("是（事後規則）" if code == "B2" else "是") if dec["回推成立"] else "否"]
    if k == 0:
        for i, v in enumerate(vals):
            cell(t10, 1, i, v)
    else:
        add_row(t10, vals)
if C is None:
    add_row(t10, ["代謝對兩者皆非、免疫對免疫", TODO, "", "", "", "", ""])
note = insert_after(x, para("註\u3000調整盛行率比為腎損傷對無腎損傷"), None)
t10._tbl.addnext(note._p)
set_text(note, "註\u3000DN 糖尿病腎病變；FSGS 局部節段性腎絲球硬化；MCD 微小變化病；TMD 薄基底膜病；MN 膜性腎病變；括號內為處理批次。"
         "95% CI 取第一次重複之折外預測；「可區分」依計算前提交之判定規則，ANCA 對 DN 一列的規則訂於已知全基因結果之後。")

# ---------- 陸、討論與柒、結論 ----------
replace_in("腎臟切片轉錄體的結果顯示", "腎臟切片轉錄體的結果顯示，免疫性病因在腎臟組織中有可辨識的分子訊號。",
           "把三個方向放在一起看，切片轉錄體在同一處理批次內，免疫性腎病可與非免疫腎病區分，免疫對代謝只有全基因模型可以區分；"
           + C_DISC + "。")
set_text(para("以 NIH 公開之腎臟切片轉錄體補充分析"),
         f"以 NIH 公開的腎臟切片轉錄體補充分析，在同一處理批次內，狼瘡腎炎與 ANCA 相關血管炎可與非免疫腎病區分（AUROC {f2(b3[P_]['auroc_mean'])}、"
         f"{f2(b1[G_]['auroc_mean'])}）" + C_CONC + "。免疫方向的訊號存在於腎臟組織的干擾素、補體與發炎細胞路徑，"
         "常規檢驗沒有反映這些路徑的項目；要回推免疫方向，仍需補體、自體抗體等專項檢驗與腎臟切片。")

# ---------- 臨床解讀（討論新增一段）----------
disc = para("把三個方向放在一起看")
insert_after(disc, disc,
  "這組結果在臨床上有兩點意義。第一，糖尿病腎病變的切片並非沒有發炎，兩批都可見補體、單核巨噬細胞與 T 細胞路徑升高，"
  "與糖尿病腎病變有先天免疫與補體活化的文獻一致（Flyvbjerg, 2017; Tesch, 2017）；以發炎路徑區分免疫與代謝方向並不可靠，"
  "ANCA 相關血管炎對糖尿病腎病變時免疫路徑模型分不開，原因可能就在這裡。第二，干擾素是目前方向性最清楚的訊號，"
  "它在狼瘡腎炎明顯升高、在糖尿病腎病變反而偏低，也能把狼瘡腎炎和同屬免疫複合體腎病的膜性腎病變分開。"
  "臨床上的推論順序因此是先以常規檢驗與病史判斷代謝與感染方向，若出現血尿、蛋白尿快速增加或腎功能短期惡化，"
  "再以補體 C3、C4、抗雙股 DNA 抗體、ANCA 與腎臟切片確認免疫方向。血液中的干擾素相關指標能否反映腎臟的干擾素訊號，需要另以資料驗證。")
replace_in("已知腎損傷時，依公開資料可先考慮代謝方向", "免疫方向三種病須以專項抗體與腎臟切片確診。",
           "免疫方向三種病須以專項抗體與腎臟切片確診；切片中干擾素訊號對狼瘡腎炎較特異，而糖尿病腎病變的切片也有補體與巨噬細胞路徑升高，"
           "以發炎指標區分免疫與代謝時須留意。")
# ---------- 摘要補上三方向 ----------
abs_zh = f"（AUROC {f2(b3[P_]['auroc_mean'])}與{f2(b1[G_]['auroc_mean'])}）。"
replace_in("腎損傷的病因大致可分為", abs_zh,
           f"（AUROC {f2(b3[P_]['auroc_mean'])}與{f2(b1[G_]['auroc_mean'])}），狼瘡腎炎也能以干擾素路徑與膜性腎病變區分"
           f"（{f2(it['glom']['auroc'])}）；糖尿病腎病變對非免疫腎病則因每批僅{lo_dn}至{hi_dn}例而未達顯著。")
abs_en = f"(AUROC {f2(b3[P_]['auroc_mean'])} and {f2(b1[G_]['auroc_mean'])})."
replace_in("Kidney damage may stem", abs_en,
           f"(AUROC {f2(b3[P_]['auroc_mean'])} and {f2(b1[G_]['auroc_mean'])}), and the interferon pathway alone separated lupus "
           f"nephritis from membranous nephropathy ({f2(it['glom']['auroc'])}). Diabetic nephropathy versus non-immune disease did not "
           f"reach significance with only {lo_dn} to {hi_dn} cases per batch.")

d.save(dst)
print("完成", dst, "（修正二未完成，含待補標記）" if C is None else "")
