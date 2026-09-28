# -*- coding: utf-8 -*-
"""科展附件與研究數據圖表（對應使用者在 Word 上修改後的作品說明書）——數字全部來自結果檔或使用者的 Word，不手打。
    python build_appendix_fair2.py <使用者說明書.docx> <輸出資料夾>
輸出：
  附件.docx／附件.md    說明書已刪除之附錄（變數字典、CDC 回推式）與補充資料（各週期樣本、免疫次樣本、通過偽發現率之暴露、資料可得性）
  研究數據表.xlsx        使用者 Word 中的全部表格（逐格照抄）＋附件各表＋暴露掃描全表＋圖二至圖九之繪圖數據
  研究數據圖/            使用者 Word 中的圖一至圖九（自 docx 取出，與文件內完全相同）＋圖一向量檔
全部內容不含版本字樣（v3、v3.1、v3.2）。Word：python -X utf8 md2docx.py <附件.md> --fair"""
import glob
import json
import os
import re
import shutil
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pandas.io.sas.sas_xport import XportReader

import build_fair_v3_2 as F

ROOT = F.ROOT
USER_DOCX, OUT = sys.argv[1], sys.argv[2]
J = lambda *p: json.load(open(os.path.join(ROOT, *p), encoding="utf-8"))
AU, IM, EXW = F.AU, F.IM, F.EXW
FIGDATA = J("figures", "fair2", "fig_data.json")
CN = "一二三四五六七八九十"
VERSION = re.compile(r"[vV]\s?3(\.\d)?|V3")


def cn(k):
    t, o = divmod(k, 10)
    return ((CN[t - 1] if t > 1 else "") + "十" if t else "") + (CN[o - 1] if o else "")


# ── 1. 使用者 Word 中的表格（逐格照抄）與圖說
def user_tables_and_captions():
    doc = Document(USER_DOCX)
    tables, figcaps, last_par, last_head = [], [], "", ""
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[1]
        if tag == "p":
            t = "".join(x.text or "" for x in child.iter(qn("w:t"))).strip()
            if not t:
                continue
            last_par = t
            if re.match(r"^[一二三四五六七八九十]+、", t):
                last_head = t
            if re.match(r"^圖[一二三四五六七八九十]+、", t):
                figcaps.append(t)
        elif tag == "tbl":
            rows = [[c.text.strip() for c in r.cells] for r in Table(child, doc).rows]
            ref = re.search(r"（(表[一二三四五六七八九十]+)）", last_par)
            if re.match(r"^表[一二三四五六七八九十]+、", last_par):
                cap = last_par
            elif ref:     # 說明書中此表沒有表號標題，但前一段內文以「（表X）」引用
                cap = f"{ref.group(1)}（說明書中此表沒有表號標題；表號依內文）"
            else:
                cap = f"{re.sub(r'^[一二三四五六七八九十]+、', '', last_head)}（說明書中此表沒有表號標題）"
            tables.append((cap, rows))
    return tables, figcaps


# ── 2. 補充資料（全部來自結果檔）
def per_cycle():
    note = {"2001-2002": "鹼性磷酸酶、LDH、磷、總膽紅素以另一變數名發布，已依官方文件對應",
            "2003-2004": "無 HCV RNA 變數：HCV 抗體陽性者只能判為未知",
            "2017-2018": "生化依 CDC 回推式換回 1999–2016 年量尺；HCV RNA 碼 3＝抗體篩檢陰性"}
    rows = [["週期", "成人", "腎臟指標 陽／陰／未知", "肝炎 陽／陰／未知", "糖尿病 陽／陰／未知", "備註"]]
    for c, r in AU["per_cycle"].items():
        k, h, d = r["kidney_v3_2"], r["hep_v3_2"], r["dm_v3_2"]
        f3 = lambda x: f"{x['pos']:,}／{x['neg']:,}／{x['unknown']:,}"
        rows.append([c.replace("-", "–"), f"{r['adults']:,}", f3(k), f3(h), f3(d),
                     note.get(c, "HCV RNA 碼 3＝抗體篩檢陰性" if c >= "2013-2014" else "")])
    tot = lambda key: {s: sum(r[key][s] for r in AU["per_cycle"].values()) for s in ("pos", "neg", "unknown")}
    k, h, d = tot("kidney_v3_2"), tot("hep_v3_2"), tot("dm_v3_2")
    assert sum(r["adults"] for r in AU["per_cycle"].values()) == AU["n_adults"] and k == {s: AU["kidney"]["v3_2"][s] for s in k}
    rows.append(["合計", f"{AU['n_adults']:,}", f"{k['pos']:,}／{k['neg']:,}／{k['unknown']:,}",
                 f"{h['pos']:,}／{h['neg']:,}／{h['unknown']:,}", f"{d['pos']:,}／{d['neg']:,}／{d['unknown']:,}", ""])
    return rows


def immune():
    by_c = [["週期", "人數", "抗核抗體 3+／4+ 陽性", "陽性比例"]]
    for c, r in IM["by_cycle"].items():
        by_c.append([c.replace("-", "–"), f"{r['n']:,}", f"{r['n_pos']}", f"{r['n_pos'] / r['n']:.1%}"])
    by_c.append(["合計", f"{IM['n']:,}", f"{IM['n_pos']}", f"{IM['prevalence']:.1%}"])
    assert sum(r["n"] for r in IM["by_cycle"].values()) == IM["n"]
    by_s = [["性別", "人數", "陽性", "陽性比例"]] + [[s, f"{r['n']:,}", f"{r['n_pos']}", f"{r['n_pos'] / r['n']:.1%}"]
                                              for s, r in IM["by_sex"].items()]
    ab = [["特異自體抗體（陽性者中）", "人數"]] + [[k, str(v)] for k, v in IM["sle_antibodies_among_pos"].items()] + \
         [["任一特異自體抗體", str(IM["any_specific_among_pos"])]]
    return by_c, by_s, ab


def exposure_labels():
    codes = {r["exposure"].replace("_percr", "") for r in EXW["results"] if not r["exposure"].startswith("藥")}
    lab = {}
    for f in sorted(glob.glob(os.path.join(ROOT, "data", "raw", "*.xpt"))):
        rd = XportReader(f)
        for fld in rd.fields:
            n = fld["name"].decode() if isinstance(fld["name"], bytes) else fld["name"]
            if n in codes and n not in lab:
                lab[n] = (fld["label"].decode("latin-1") if isinstance(fld["label"], bytes) else fld["label"]).strip()
        rd.close()
    assert codes <= set(lab), sorted(codes - set(lab))
    return lab


LAYER = {"M0_粗關聯": "未調整", "M1_人口學": "人口學", "M2_＋體位抽菸": "人口學、體位、抽菸",
         "M3_＋糖尿病高血壓": "人口學、體位、抽菸、糖尿病、高血壓", "M4_藥物＋總用藥數": "人口學、體位、抽菸、糖尿病、高血壓、總用藥數"}
ZH = {"LBXBCD": "血鎘", "LBXBPB": "血鉛", "LBXTHG": "血汞", "LBXBSE": "血硒", "URXUTL": "尿鉈", "URXUCS": "尿銫", "URXUSR": "尿鍶",
      "URXUBA": "尿鋇", "URXUSB": "尿銻", "URXUCD": "尿鎘", "URXUPB": "尿鉛", "URXUAS": "尿總砷", "LBXPFHS": "血清 PFHxS",
      "LBXPFOA": "血清 PFOA", "LBXNFOA": "血清 n-PFOA"}     # 其餘以 NHANES 官方英文標籤呈現


def exposure_rows(lab):
    """暴露掃描全表：主要結果為事前寫定之調整模型（藥物＋總用藥數，或糖尿病高血壓），q 值以該層之 p 值計算。"""
    def name(e):
        if e.startswith("藥陰_"):
            return f"{e[3:]}（陰性對照用藥）"
        if e.startswith("藥_"):
            return e[2:].replace("_", "／")
        code = e.replace("_percr", "")
        return (ZH.get(code) or lab[code]) + ("（肌酸酐比值）" if e.endswith("_percr") else "")

    def layer(r):
        h = r["headline"]
        ks = [k for k, m in r["models"].items() if m["OR"] == h["OR"] and m["n"] == h["n"]]
        assert len(ks) == 1, (r["exposure"], ks)
        return LAYER[ks[0]]
    rows = [["暴露", "NHANES 代碼", "NHANES 官方標籤", "量尺", "主要模型之調整變項", "n", "腎臟指標異常人數", "勝算比", "95% CI 下限",
             "95% CI 上限", "p 值", "q 值（偽發現率）", "通過偽發現率 0.05", "通過完整調整", "反向因果風險"]]
    for r in sorted(EXW["results"], key=lambda r: r["q_bh"]):
        h, e = r["headline"], r["exposure"]
        code = "" if e.startswith("藥") else e.replace("_percr", "")
        rows.append([name(e), code, lab.get(code, ""), h["scale"], layer(r), h["n"], h["n_pos"], round(h["OR"], 4),
                     round(h["ci"][0], 4), round(h["ci"][1], 4), float(f"{h['p_two_sided']:.3g}"), float(f"{r['q_bh']:.3g}"),
                     "是" if r["significant_fdr05"] else "否", "是" if r["survives_full_adjustment"] else "否", r["reverse_causation_risk"]])
    assert len(rows) - 1 == EXW["n_scanned"] and sum(x[12] == "是" for x in rows[1:]) == EXW["n_significant_fdr05"]
    return rows


# ── 3. 圖二至圖九之繪圖數據（與 figures/fair2/fig_data.json 相同，攤平成表）
def fig_data_sheets():
    D, out = FIGDATA, []
    b = D["圖二"]
    out.append(("圖二數據", [["方框", "文字"], ["最上方", b["top_box"]], ["腎臟指標", b["kidney_box"]]] +
                [[x["title"], x["text"]] for x in b["label_boxes"]] + [["說明", b["note"]], ["資料修正", b["fix_box_title"]]] +
                [["", t] for t in b["fix_box_items"]]))
    rows = [["方向", "指標", "模型", "五次重複平均", "最小", "最大"]]
    for ax, d in D["圖三"].items():
        for met in ("auroc", "ap"):
            rows += [[d["title"], "AUROC" if met == "auroc" else "平均精確率（AP）", r["label"], r["mean"], r["min"], r["max"]] for r in d[met]]
    out.append(("圖三數據", rows))
    rows = [["方向", "模型", "類別", "項目", "數值 1", "數值 2"]]
    for ax, d in D["圖四"].items():
        for m in d["models"]:
            rows.append([d["title"], m["label"], "校準", "截距／斜率", m["intercept"], m["slope"]])
            rows += [[d["title"], m["label"], "校準曲線", f"第 {i} 點：平均預測／實際陽性比例", p["mean_pred"], p["observed"]]
                     for i, p in enumerate(m["curve"], 1)]
            rows += [[d["title"], m["label"], "三段分區", f"{k}：實際陽性比例／人數", v["observed_rate"], v["n"]] for k, v in m["bands"].items()]
        rows.append([d["title"], "", "盛行率", "", d["prevalence"], ""])
    out.append(("圖四數據", rows))
    rows = [["方向", "模型", "條件", "AUROC", "95% CI 下限", "95% CI 上限"]]
    for ax, d in D["圖五"].items():
        rows += [[d["title"], m["label"], c["label"], c["auroc"], *(c["ci95"] or ["", ""])] for m in d["models"] for c in m["conditions"]]
    out.append(("圖五數據", rows))
    rows = [["方向", "組別", "陽性／人數", "模型", "AUROC", "95% CI 下限", "95% CI 上限"]]
    for ax, d in D["圖六"].items():
        rows += [[d["title"], g["name"], f"{g['n_pos']}／{g['n']:,}", r["label"], r["auroc"], *r["ci95"]] for g in d["groups"] for r in g["rows"]]
    out.append(("圖六數據", rows))
    rows = [["方向", "閾值機率"] + [c["label"] for c in D["圖七"]["肝炎"]["curves"]]]
    for ax, d in D["圖七"].items():
        rows += [[d["title"], pt] + [c["net_benefit"][i] for c in d["curves"]] for i, pt in enumerate(d["pt"])]
        rows.append([d["title"], "註記", d["annotation"]["text"].replace("\n", "")] + [""] * (len(d["curves"]) - 1))
    out.append(("圖七數據", rows))
    rows = [["金屬", "同一批人數", "腎損傷定義", "測量", "勝算比（濃度加倍）", "95% CI 下限", "95% CI 上限"]]
    for r in D["圖八"]["rows"]:
        rows += [[r["metal"], r["n_both"], col, e["label"], e["or"], *e["ci95"]] for col, panel in zip(D["圖八"]["columns"], r["panels"]) for e in panel]
    out.append(("圖八數據", rows))
    rows = [["方向", "測試半陽性中位數", "做法", "平均預測／實際：中位數", "2.5 百分位", "97.5 百分位"]]
    for ax, d in D["圖九"].items():
        rows += [[d["title"], d["n_pos_test_median"], r["label"], r["median"], r["p2_5"], r["p97_5"]] for r in d["rows"]]
    out.append(("圖九數據", rows))
    return out


# ── 4. 附件 Word（markdown → md2docx --fair）
def md_table(rows):
    esc = lambda s: str(s).replace("|", "／")
    return "\n".join(["| " + " | ".join(esc(c) for c in rows[0]) + " |", "|" + "---|" * len(rows[0])] +
                     ["| " + " | ".join(esc(c) for c in r) + " |" for r in rows[1:]])


def appendix_md(cyc, imm, sig_rows):
    by_c, by_s, ab = imm
    conv_head = ["變數", "名稱", "2017–2018 → 1999–2016 年量尺（BIOPRO_J）", "2021–2023 → 2017–2020 年量尺"]
    k = 0

    def sec(title):
        nonlocal k
        k += 1
        return f"## 附件{cn(k)}、{title}"
    return f"""# {F.TITLE}

## ——作品說明書附件

本附件收錄作品說明書未列出的補充資料。說明書中的圖一至圖九另以原始圖檔（「研究數據圖」資料夾）提供；表一至表十八、本附件各表、上游暴露掃描全表（{EXW['n_scanned']} 項），以及圖二至圖九的繪圖數據，都收錄於「研究數據表.xlsx」。所有數字均由分析結果檔產生。

{sec("常規套組變數字典")}

「近端排除」表示該軸因生理上緊鄰標籤而移除（例如肝炎軸移除肝功能指標、糖尿病軸移除血糖）。

| 代號 | 名稱 | 單位 | 肝炎軸 | 糖尿病軸 |
|---|---|---|---|---|
{F.DICT_ROWS}

{sec("CDC 官方回推式（X 為換算後、Y 為原值）")}

{md_table([conv_head])}
{F.CONV_ROWS}

{sec("各週期分析樣本與標籤")}

腎臟指標為 eGFR < 60 或 ACR ≧ 30 之三值判定；肝炎與糖尿病標籤只在腎臟指標異常者中判定。

{md_table(cyc)}

{sec("免疫方向：抗核抗體次樣本")}

抗核抗體只在 1999–2004 年剩餘血清次樣本檢驗；以下為腎臟指標異常且在次樣本中的成人。

{md_table(by_c)}

{md_table(by_s)}

{md_table(ab)}

{sec(f"上游暴露掃描：通過偽發現率 0.05 之 {EXW['n_significant_fdr05']} 項暴露")}

以全體成人掃描 {EXW['n_scanned']} 個暴露（Benjamini–Hochberg 法控制偽發現率）；依 q 值排序。這些關聯分不出先後，不能當作致病證據（見說明書肆之六）；全部 {EXW['n_scanned']} 項見研究數據表.xlsx。

{md_table(sig_rows)}

{sec("資料與程式可得性")}

本研究只使用美國國家健康與營養調查（NHANES）公開、已去識別的資料檔，共 {F.N_ALL} 個檔案（開發資料 {F.N_DEV} 個、外部資料 {F.N_EXT} 個），每個檔案的來源網址、SHA256 雜湊與位元組數都記錄在出處帳本。分析程式與結果存放於 Git 儲存庫（目錄 experiments/kidney_cause）；每一項分析都先提交分析計畫、再執行，提交代碼見說明書表七。

> ⚠️ **待確認**：是否附上儲存庫網址。網址與提交紀錄可能透露作者或學校，請依比賽規定決定。
"""


# ── 5. Excel
def write_xlsx(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    toc = wb.create_sheet("目錄")
    toc.append(["工作表", "內容"])
    bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="D9D9D9")
    for name, title, rows in sheets:
        toc.append([name, title])
        ws = wb.create_sheet(name)
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([])
        for i, r in enumerate(rows):
            ws.append(r)
            if i == 0:
                for c in ws[ws.max_row]:
                    c.font, c.fill = bold, head_fill
        width = {}
        for r in rows:
            for j, v in enumerate(r):
                s = str(v)
                width[j] = max(width.get(j, 0), sum(2 if ord(ch) > 0x2E7F else 1 for ch in s))
        for j, w in width.items():
            ws.column_dimensions[ws.cell(row=3, column=j + 1).column_letter].width = min(max(w + 2, 8), 60)
        for row in ws.iter_rows(min_row=3):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A4"
    for c in toc[1]:
        c.font, c.fill = bold, head_fill
    toc.column_dimensions["A"].width, toc.column_dimensions["B"].width = 18, 90
    wb.save(path)


def main():
    os.makedirs(os.path.join(OUT, "研究數據圖"), exist_ok=True)
    tables, figcaps = user_tables_and_captions()
    assert len(figcaps) == 9, figcaps
    # 研究數據圖：自使用者 docx 取出圖一至圖九（依出現順序），檔名用圖說
    z = zipfile.ZipFile(USER_DOCX)
    rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
    rid = {a["Id"]: a["Target"] for a in (dict(re.findall(r'(\w+)="([^"]*)"', m)) for m in re.findall(r"<Relationship [^>]*>", rels))}
    embeds = [rid[r] for r in re.findall(r'r:embed="(rId\d+)"', z.read("word/document.xml").decode("utf-8"))]
    assert len(embeds) == 9
    for cap, target in zip(figcaps, embeds):
        fname = re.sub(r'[\\/:*?"<>|]', "", cap.replace("、", "_", 1)).replace(" ", "")[:40] + ".png"
        open(os.path.join(OUT, "研究數據圖", fname), "wb").write(z.read("word/" + target))
    shutil.copy2(os.path.join(ROOT, "figures", "fair", "圖一_研究架構示意圖.svg"), os.path.join(OUT, "研究數據圖", "圖一_研究架構示意圖（向量檔）.svg"))
    # 附件資料
    cyc, imm = per_cycle(), immune()
    lab = exposure_labels()
    exw = exposure_rows(lab)
    sig = [["暴露", "勝算比（95% CI）", "量尺", "q 值", "n"]] + \
          [[r[0], f"{r[7]:.3f}（{r[8]:.3f}–{r[9]:.3f}）", r[3], f"{r[11]:.1e}", f"{r[5]:,}"] for r in exw[1:] if r[12] == "是"]
    md = appendix_md(cyc, imm, sig)
    assert not VERSION.search(md), VERSION.search(md)
    open(os.path.join(OUT, "附件.md"), "w", encoding="utf-8").write(md)
    # Excel：使用者表格 → 附件各表 → 暴露全表 → 繪圖數據
    sheets = []
    for i, (cap, rows) in enumerate(tables, 1):
        m = re.match(r"^(表[一二三四五六七八九十]+)", cap)
        sheets.append((m.group(1) if m else cap.split("（")[0][:20], cap, rows))
    dict_rows = [["代號", "名稱", "單位", "肝炎軸", "糖尿病軸"]] + [r.strip("| ").split(" | ") for r in F.DICT_ROWS.split("\n")]
    conv_rows = [["變數", "名稱", "2017–2018 → 1999–2016 年量尺（BIOPRO_J）", "2021–2023 → 2017–2020 年量尺"]] + \
                [r.strip("| ").split(" | ") for r in F.CONV_ROWS.split("\n")]
    by_c, by_s, ab = imm
    sheets += [("附件一_變數字典", "附件一、常規套組變數字典", dict_rows),
               ("附件二_回推式", "附件二、CDC 官方回推式（X 為換算後、Y 為原值）", conv_rows),
               ("附件三_各週期", "附件三、各週期分析樣本與標籤", cyc),
               ("附件四_免疫次樣本", "附件四、免疫方向：抗核抗體次樣本", by_c + [[]] + by_s + [[]] + ab),
               ("附件五_暴露掃描全表", f"附件五、上游暴露掃描全表（{EXW['n_scanned']} 項，依 q 值排序；主要調整模型之勝算比）", exw)]
    sheets += [(n, f"{n[:2]}之繪圖數據（與說明書{n[:2]}相同）", rows) for n, rows in fig_data_sheets()]
    for name, title, rows in sheets:
        assert not VERSION.search(title + json.dumps(rows, ensure_ascii=False)), (name, VERSION.search(json.dumps(rows, ensure_ascii=False)))
    write_xlsx(os.path.join(OUT, "研究數據表.xlsx"), sheets)
    print(f"[完成] {OUT}：表 {len(tables)} 個（使用者 Word）＋附件表 5 個＋繪圖數據 8 張；圖 {len(figcaps)} 張；暴露 {len(exw) - 1} 項（通過 {len(sig) - 1}）")


if __name__ == "__main__":
    main()
